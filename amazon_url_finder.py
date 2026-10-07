from __future__ import annotations

import html
import json
import re
import time
from pathlib import Path
from urllib.parse import urlparse

ASIN_RE = re.compile(r"[A-Z0-9]{10}")
IMAGE_RE = re.compile(r"https://m\.media-amazon\.com/images/I/[^\"'<>\s\\]+", re.IGNORECASE)


def read_asins(source_file: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(source_file, read_only=True, data_only=True)
    try:
        asins: list[str] = []
        for row_number, (value,) in enumerate(workbook.active.iter_rows(min_col=1, max_col=1, values_only=True), start=1):
            if value is None or not str(value).strip():
                continue
            asin = str(value).strip().upper()
            if row_number == 1 and asin in {"ASIN", "ASINS"}:
                continue
            if not ASIN_RE.fullmatch(asin):
                raise ValueError(f"Row {row_number}: invalid ASIN {asin!r}. Expected 10 letters or digits in column A.")
            asins.append(asin)
            if len(asins) > 5000:
                raise ValueError("The Excel file has more than 5000 ASINs.")
        if not asins:
            raise ValueError("No ASINs found in column A.")
        return asins
    finally:
        workbook.close()


def normalize_image_url(value: str) -> str:
    url = html.unescape(value.replace("\\/", "/").replace("\\u0026", "&"))
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.netloc.lower() != "m.media-amazon.com" or not parsed.path.startswith("/images/I/"):
        return "NOT FOUND"
    return re.sub(r"\._[^.]*_\.", "._AC_SL1500_.", url)


def image_from_html(source: str) -> str:
    cleaned = html.unescape(source.replace("\\/", "/").replace("\\u0026", "&"))
    for label in ("hiRes", "large", "mainUrl"):
        match = re.search(r'"' + label + r'"\s*:\s*"(https://m\.media-amazon\.com/images/I/[^\"]+)"', cleaned)
        if match:
            url = normalize_image_url(match.group(1))
            if url != "NOT FOUND":
                return url
    match = re.search(r'<meta[^>]+property=["\']og:image["\'][^>]+content=["\']([^"\']+)', cleaned, re.IGNORECASE)
    if match:
        url = normalize_image_url(match.group(1))
        if url != "NOT FOUND":
            return url
    for match in IMAGE_RE.finditer(cleaned):
        url = normalize_image_url(match.group(0))
        if url != "NOT FOUND":
            return url
    return "NOT FOUND"


def image_from_element(driver: object) -> str:
    from selenium.webdriver.common.by import By

    for by, selector in ((By.ID, "landingImage"), (By.CSS_SELECTOR, "#imgTagWrapperId img"),
                         (By.CSS_SELECTOR, "#main-image-container img"), (By.CSS_SELECTOR, "img.a-dynamic-image")):
        try:
            image = driver.find_element(by, selector)
            data = image.get_attribute("data-a-dynamic-image")
            if data:
                choices = json.loads(data)
                if isinstance(choices, dict) and choices:
                    biggest = max(choices.items(), key=lambda pair: pair[1][0] * pair[1][1])[0]
                    url = normalize_image_url(biggest)
                    if url != "NOT FOUND":
                        return url
            url = normalize_image_url(image.get_attribute("src") or "")
            if url != "NOT FOUND":
                return url
        except Exception:
            continue
    return "NOT FOUND"


def create_driver():
    from selenium import webdriver

    options = webdriver.ChromeOptions()
    options.page_load_strategy = "eager"
    for argument in ("--headless=new", "--disable-gpu", "--no-sandbox", "--disable-dev-shm-usage",
                     "--disable-popup-blocking", "--disable-extensions", "--no-first-run",
                     "--no-default-browser-check", "--lang=de-DE", "--window-size=1920,1080"):
        options.add_argument(argument)
    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(25)
    return driver


def lookup_image(driver: object, asin: str) -> str:
    try:
        driver.get(f"https://www.amazon.de/dp/{asin}")
    except Exception:
        pass
    for _ in range(3):
        try:
            url = image_from_html(driver.page_source)
            if url != "NOT FOUND":
                return url
            url = image_from_element(driver)
            if url != "NOT FOUND":
                return url
        except Exception:
            pass
        time.sleep(0.5)
    return "NOT FOUND"


def find_images(asins: list[str], output_path: Path, progress) -> tuple[int, int]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ASIN Image Links"
    sheet.append(["ASIN", "Image Link"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 90
    found = 0
    driver = create_driver()
    try:
        for index, asin in enumerate(asins, start=1):
            image_url = lookup_image(driver, asin)
            if image_url != "NOT FOUND":
                found += 1
            sheet.append([asin, image_url])
            if image_url != "NOT FOUND":
                sheet.cell(index + 1, 2).hyperlink = image_url
            progress(index, len(asins), asin, image_url)
    finally:
        driver.quit()
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output_path)
    return found, len(asins) - found
