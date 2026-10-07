from __future__ import annotations

import html
import cgi
from decimal import Decimal, InvalidOperation
from concurrent.futures import ThreadPoolExecutor, as_completed
import hashlib
import hmac
import json
import math
import os
import re
import secrets
import shutil
import subprocess
import threading
import time
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

try:
    from PIL import Image, ImageChops, ImageOps
except Exception:
    Image = None
    ImageChops = None
    ImageOps = None

try:
    import numpy as np
except Exception:
    np = None


VS_ROOT = Path(os.environ.get("VS_ROOT", "/var/www/primexeu.com.shared/photos/VS")).resolve()
APP_ROOT = Path(__file__).resolve().parent
ASSET_ROOT = Path(os.environ.get("VS_DASHBOARD_ASSET_ROOT", str(APP_ROOT / "static"))).resolve()
USER_HOME = Path(os.environ.get("USERPROFILE") or os.environ.get("HOME") or str(APP_ROOT))
BASE_URL = os.environ.get("VS_BASE_URL", "https://primexeu.com/photos/VS/").rstrip("/") + "/"
HOST = os.environ.get("VS_DASHBOARD_HOST", "127.0.0.1")
PORT = int(os.environ.get("VS_DASHBOARD_PORT", "8791"))
LOGIN_EMAIL = os.environ.get("VS_DASHBOARD_USER", "info@primexeu.com")
LOGIN_PASSWORD = os.environ.get("VS_DASHBOARD_PASSWORD", "")
SESSION_SECRET = os.environ.get("VS_DASHBOARD_SESSION_SECRET") or secrets.token_urlsafe(32)
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
URL_EXTENSIONS = IMAGE_EXTENSIONS | {".pdf"}
LOGO_SCAN_EXTENSIONS = IMAGE_EXTENSIONS | {".gif", ".bmp", ".tif", ".tiff", ".pjg"}
YEAR_RE = re.compile(r"\b(?:16|17|18|19|20)\d{2}\b")
RIEGEL_RE = re.compile(r"r\s*[i1l]\s*[e3]\s*g\s*[e3]\s*l", re.IGNORECASE)
RIEGEL_REFERENCE_PATHS = [
    USER_HOME / "Downloads" / "31s3VfcxY4L._AC_SL1500_.jpg",
    Path(r"W:\05_CLIENTS\02_VS\10_DESIGN\01_WINE-SPIRITS\25_Riegel_V_PROJECT\ALL IMAGES"),
]
RIEGEL_REFERENCE_URLS = [
    "https://m.media-amazon.com/images/I/31s3VfcxY4L._AC_SL1500_.jpg",
]
GENERATED_NAMES = {"1", "3", "6", "12", "12-alt", "12_alt"}
GRID_LAYOUTS = {"1": (1, 1), "3": (3, 1), "6": (3, 2), "12": (6, 2), "12_alt": (4, 3)}
GRID_OUTER_PADDING = 40
GRID_ITEM_GAP = 40
SCALE_BOOST = {"1": 1.35, "3": 1.18, "6": 1.08, "12": 1.0, "12_alt": 1.0}
JOBS: dict[str, dict[str, object]] = {}
JOBS_LOCK = threading.Lock()


def sign_session(value: str) -> str:
    return hmac.new(SESSION_SECRET.encode("utf-8"), value.encode("utf-8"), hashlib.sha256).hexdigest()


def make_session_cookie() -> str:
    value = f"{LOGIN_EMAIL}|{int(time.time())}|{secrets.token_urlsafe(18)}"
    return f"{value}|{sign_session(value)}"


def parse_cookies(header: str) -> dict[str, str]:
    cookies = {}
    for part in header.split(";"):
        if "=" in part:
            name, value = part.strip().split("=", 1)
            cookies[name] = value
    return cookies


def is_valid_session(cookie_value: str) -> bool:
    parts = cookie_value.split("|")
    if len(parts) != 4:
        return False
    value = "|".join(parts[:3])
    if parts[0] != LOGIN_EMAIL:
        return False
    return hmac.compare_digest(parts[3], sign_session(value))


def is_safe_name(value: str) -> bool:
    return bool(value) and "/" not in value and "\\" not in value and value not in {".", ".."}


def is_under_root(path: Path) -> bool:
    try:
        path.resolve().relative_to(VS_ROOT)
        return True
    except ValueError:
        return False


def stat_payload(path: Path) -> dict[str, object]:
    stat = path.stat()
    return {
        "name": path.name,
        "modified": int(stat.st_mtime),
        "modifiedText": time.strftime("%Y-%m-%d %H:%M", time.localtime(stat.st_mtime)),
    }


def top_path(top_name: str) -> Path:
    if not is_safe_name(top_name):
        raise ValueError("Select a valid top folder.")
    path = (VS_ROOT / top_name).resolve()
    if not path.exists() or not path.is_dir() or not is_under_root(path):
        raise ValueError("Selected top folder does not exist.")
    return path


def product_path(top_name: str, product_name: str) -> Path:
    top = top_path(top_name)
    if not is_safe_name(product_name):
        raise ValueError("Select a valid product folder.")
    path = (top / product_name).resolve()
    if not path.exists() or not path.is_dir() or not is_under_root(path):
        raise ValueError("Selected product folder does not exist.")
    return path


def count_images(folder: Path) -> int:
    return sum(1 for path in folder.rglob("*") if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS)


def list_top_folders() -> list[dict[str, object]]:
    folders = []
    for path in VS_ROOT.iterdir():
        if path.is_dir():
            item = stat_payload(path)
            item["subfolderCount"] = sum(1 for child in path.iterdir() if child.is_dir())
            folders.append(item)
    folders.sort(key=lambda item: (-int(item["modified"]), str(item["name"]).lower()))
    return folders


def list_product_folders(top_name: str) -> list[dict[str, object]]:
    try:
        top = top_path(top_name)
    except ValueError:
        return []
    products = []
    for path in top.iterdir():
        if path.is_dir():
            item = stat_payload(path)
            item["imageCount"] = count_images(path)
            products.append(item)
    products.sort(key=lambda item: (-int(item["modified"]), str(item["name"]).lower()))
    return products


def public_url_for(path: Path, cache_bust: bool = False) -> str:
    encoded = "/".join(quote(part) for part in path.resolve().relative_to(VS_ROOT).parts)
    url = BASE_URL + encoded
    if not cache_bust:
        return url
    try:
        return f"{url}?v={int(path.stat().st_mtime)}"
    except OSError:
        return url


def check_public_url(url: str, timeout: int = 8) -> tuple[bool, str]:
    try:
        request = Request(url, method="HEAD", headers={"User-Agent": "VS-Tools-URL-Check/1.0"})
        with urlopen(request, timeout=timeout) as response:
            status = response.status
    except HTTPError as exc:
        if exc.code not in {403, 405}:
            return False, f"ERROR {exc.code}"
        try:
            request = Request(url, method="GET", headers={"User-Agent": "VS-Tools-URL-Check/1.0", "Range": "bytes=0-0"})
            with urlopen(request, timeout=timeout) as response:
                status = response.status
        except HTTPError as get_exc:
            return False, f"ERROR {get_exc.code}"
        except URLError as get_exc:
            return False, f"ERROR {get_exc.reason}"
        except Exception as get_exc:
            return False, f"ERROR {get_exc}"
    except URLError as exc:
        return False, f"ERROR {exc.reason}"
    except Exception as exc:
        return False, f"ERROR {exc}"
    return (200 <= status < 400), ("OK" if 200 <= status < 400 else f"ERROR {status}")


def image_files(folder: Path, recursive: bool, include_pdf: bool = False) -> list[Path]:
    allowed = URL_EXTENSIONS if include_pdf else IMAGE_EXTENSIONS
    files = folder.rglob("*") if recursive else folder.iterdir()
    return sorted([p for p in files if p.is_file() and p.suffix.lower() in allowed], key=lambda p: str(p).lower())


def generate_urls(top_name: str, product_name: str = "", recursive: bool = True, cache_bust: bool = False) -> list[dict[str, str]]:
    target = product_path(top_name, product_name) if product_name else top_path(top_name)
    rows = []
    for file_path in image_files(target, recursive, include_pdf=True):
        rel = str(file_path.relative_to(VS_ROOT)).replace(os.sep, "/")
        rows.append({
            "folder": top_name,
            "product": product_name or (rel.split("/")[1] if "/" in rel else ""),
            "file": file_path.name,
            "path": rel,
            "url": public_url_for(file_path, cache_bust),
        })
    return rows


def autocrop_whitespace(img, tol: int = 10, alpha_tol: int = 2):
    if img.mode != "RGBA":
        img = img.convert("RGBA")
    alpha = img.split()[-1]
    if alpha.getextrema()[0] < 255:
        alpha_mask = alpha.point(lambda p: 255 if p > alpha_tol else 0)
        bbox = alpha_mask.getbbox()
        if bbox:
            return img.crop(bbox)
    bg = Image.new("RGBA", img.size, (255, 255, 255, 255))
    diff = ImageChops.difference(img, bg)
    diff = ImageOps.grayscale(diff)
    mask = diff.point(lambda p: 255 if p > tol else 0)
    bbox = mask.getbbox()
    return img.crop(bbox) if bbox else img


def resize_rgba_alpha_safe(img, size, resample=Image.Resampling.LANCZOS):
    if img.mode != "RGBA":
        img = img.convert("RGBA")
    alpha = img.getchannel("A")
    if alpha.getextrema() == (255, 255):
        return img.resize(size, resample)
    return img.convert("RGBa").resize(size, resample).convert("RGBA")


def flatten_rgba_on_white(img):
    if img.mode != "RGBA":
        img = img.convert("RGBA")
    white = Image.new("RGBA", img.size, (255, 255, 255, 255))
    white.alpha_composite(img)
    return white.convert("RGB")


def generate_composited_grid(source_img, quantity_key: str, canvas_size: int = 2500, spacing: int = GRID_ITEM_GAP):
    cols, rows = GRID_LAYOUTS[quantity_key]
    src_rgba = autocrop_whitespace(source_img.convert("RGBA"))
    max_item_w = (canvas_size - (GRID_OUTER_PADDING * 2) - spacing * (cols - 1)) / cols
    max_item_h = (canvas_size - (GRID_OUTER_PADDING * 2) - spacing * (rows - 1)) / rows
    scale = min(max_item_w / src_rgba.width, max_item_h / src_rgba.height) * SCALE_BOOST.get(quantity_key, 1.0)
    scale = min(scale, max_item_w / src_rgba.width, max_item_h / src_rgba.height)
    new_size = (max(1, int(src_rgba.width * scale)), max(1, int(src_rgba.height * scale)))
    resized = resize_rgba_alpha_safe(src_rgba, new_size, Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (canvas_size, canvas_size), (255, 255, 255, 0))
    group_w = cols * resized.width + (cols - 1) * spacing
    group_h = rows * resized.height + (rows - 1) * spacing
    start_x = int((canvas_size - group_w) / 2)
    start_y = int((canvas_size - group_h) / 2)
    for row in range(rows):
        for col in range(cols):
            x = int(start_x + col * (resized.width + spacing))
            y = int(start_y + row * (resized.height + spacing))
            canvas.alpha_composite(resized, (x, y))
    return canvas


def source_candidates(folder: Path) -> list[Path]:
    candidates = []
    for path in folder.iterdir():
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS and path.stem.lower() not in GENERATED_NAMES:
            candidates.append(path)
    return sorted(candidates, key=lambda p: (p.stat().st_mtime, p.name.lower()))


def target_product_folders(top_name: str, product_name: str) -> list[Path]:
    if product_name:
        return [product_path(top_name, product_name)]
    return [VS_ROOT / top_name / item["name"] for item in list_product_folders(top_name)]


def generate_grid_images(
    top_name: str,
    product_name: str,
    quantities: list[str],
    export_format: str,
    delete_source: bool,
    cache_bust: bool = False,
) -> list[dict[str, str]]:
    if Image is None:
        raise ValueError("Pillow is not installed on the server.")
    quantities = [q for q in quantities if q in GRID_LAYOUTS]
    if not quantities:
        raise ValueError("Select at least one grid quantity.")
    if export_format not in {"jpg", "png"}:
        raise ValueError("Select JPG or PNG.")
    rows = []
    for folder in target_product_folders(top_name, product_name):
        candidates = source_candidates(folder)
        if not candidates:
            rows.append({"product": folder.name, "status": "skipped", "message": "No source image found"})
            continue
        source = candidates[0]
        saved = 0
        try:
            with Image.open(source) as img:
                source_img = img.convert("RGBA")
                for quantity in quantities:
                    out_name = f"{quantity.replace('_alt', '-alt')}.{export_format}"
                    out_path = folder / out_name
                    result = generate_composited_grid(source_img, quantity)
                    if export_format == "jpg":
                        flatten_rgba_on_white(result).save(out_path, quality=95, optimize=True)
                    else:
                        result.save(out_path, optimize=True)
                    saved += 1
                    rows.append({
                        "product": folder.name,
                        "source": source.name,
                        "created": out_path.name,
                        "url": public_url_for(out_path, cache_bust),
                        "status": "created",
                    })
            if delete_source and saved == len(quantities):
                source.unlink()
        except Exception as exc:
            rows.append({"product": folder.name, "source": source.name, "status": "error", "message": str(exc)})
    return rows


def job_snapshot(job_id: str) -> dict[str, object] | None:
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            return None
        return {
            "id": job_id,
            "done": job.get("done", False),
            "ok": job.get("ok", True),
            "phase": job.get("phase", ""),
            "message": job.get("message", ""),
            "percent": job.get("percent", 0),
            "uploaded": job.get("uploaded", 0),
            "created": job.get("created", 0),
            "urls": job.get("urls", 0),
            "years": job.get("years", 0),
            "scanned": job.get("scanned", 0),
            "found": job.get("found", 0),
            "issues": job.get("issues", 0),
            "rows": list(job.get("rows", [])),
        }


def update_job(job_id: str, **values: object) -> None:
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if job:
            job.update(values)


def append_job_rows(job_id: str, rows: list[dict[str, str]]) -> None:
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if job:
            job_rows = job.setdefault("rows", [])
            if isinstance(job_rows, list):
                job_rows.extend(rows)


def finish_job_counts(job_id: str) -> None:
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            return
        rows = job.get("rows", [])
        if not isinstance(rows, list):
            rows = []
        successful = sum(1 for row in rows if isinstance(row, dict) and row.get("status") in {"uploaded", "uploaded then organized", "moved", "copied", "created", "ready", "year found", "no year"})
        job["created"] = sum(1 for row in rows if isinstance(row, dict) and row.get("step") == "grid" and row.get("status") == "created")
        job["urls"] = sum(1 for row in rows if isinstance(row, dict) and row.get("step") == "url" and row.get("url"))
        job["years"] = sum(1 for row in rows if isinstance(row, dict) and row.get("status") == "year found")
        job["issues"] = len(rows) - successful


def run_upload_automation_job(
    job_id: str,
    relative_parts: list[str],
    quantities: list[str],
    export_format: str,
    delete_source: bool,
    cache_bust: bool = False,
) -> None:
    try:
        top_name = relative_parts[0]
        product_name = relative_parts[1] if len(relative_parts) > 1 else ""
        products_for_grid: list[str] = []
        final_url_rows: list[dict[str, str]] = []

        update_job(job_id, phase="organizing", percent=12, message="Organizing uploaded photos into product folders...")
        if product_name:
            append_job_rows(job_id, [{"step": "organize", "product": product_name, "status": "ready", "message": "Already uploaded into a product folder"}])
            products_for_grid = [product_name]
        else:
            organize_rows = organize_photos(top_name, top_name, True)
            for row in organize_rows:
                row["step"] = "organize"
                if delete_source and row.get("status") == "moved":
                    row.pop("url", None)
                    row["message"] = "Temporary source; removed after successful grid generation"
            with JOBS_LOCK:
                job = JOBS.get(job_id)
                if job and isinstance(job.get("rows"), list):
                    for row in job["rows"]:
                        if isinstance(row, dict) and row.get("status") == "uploaded":
                            row.pop("url", None)
                            row["status"] = "uploaded then organized"
                            row["message"] = "Moved into product folder during automation"
                    job["rows"].extend(organize_rows)
            products_for_grid = sorted({row.get("product", "") for row in organize_rows if row.get("status") == "moved" and row.get("product")})

        total_products = max(1, len(products_for_grid))
        for index, product in enumerate(products_for_grid, start=1):
            update_job(
                job_id,
                phase="generating grids",
                percent=20 + int((index - 1) / total_products * 65),
                message=f"Generating grid images for {product} ({index}/{len(products_for_grid)})...",
            )
            grid_rows = generate_grid_images(top_name, product, quantities, export_format, delete_source, cache_bust)
            for row in grid_rows:
                row["step"] = "grid"
                if row.get("status") == "created" and row.get("url"):
                    final_url_rows.append({
                        "step": "url",
                        "product": row.get("product", product),
                        "file": row.get("created", ""),
                        "status": "ready",
                        "url": row["url"],
                    })
            append_job_rows(job_id, grid_rows)
            finish_job_counts(job_id)

        update_job(job_id, phase="collecting urls", percent=92, message="Collecting final URLs...")
        append_job_rows(job_id, final_url_rows)
        finish_job_counts(job_id)
        update_job(job_id, done=True, ok=True, phase="done", percent=100, message="Finished.")
    except Exception as exc:
        append_job_rows(job_id, [{"step": "automation", "status": "error", "message": str(exc)}])
        finish_job_counts(job_id)
        update_job(job_id, done=True, ok=False, phase="error", percent=100, message=f"Automation failed: {exc}")


def organize_photos(source_top: str, destination_top: str, move_files: bool) -> list[dict[str, str]]:
    source = top_path(source_top)
    if not is_safe_name(destination_top):
        raise ValueError("Enter or select a valid destination folder.")
    destination = (VS_ROOT / destination_top).resolve()
    if not is_under_root(destination):
        raise ValueError("Destination must stay inside the VS folder.")
    destination.mkdir(parents=True, exist_ok=True)

    rows = []
    for file_path in sorted(source.iterdir(), key=lambda p: p.name.lower()):
        if not file_path.is_file() or file_path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        code = re.split(r"[_\\-\\s]+", file_path.stem.strip(), maxsplit=1)[0]
        code = "".join(char for char in code if char.isalnum())
        if code.lower().endswith("bio") and code[:-3].isdigit():
            code = code[:-3]
        if not code:
            rows.append({"file": file_path.name, "status": "skipped", "message": "Could not detect product code from filename"})
            continue
        product_folder = destination / code
        product_folder.mkdir(exist_ok=True)
        target = product_folder / file_path.name
        if move_files:
            file_path.replace(target)
            action = "moved"
        else:
            import shutil
            shutil.copy2(file_path, target)
            action = "copied"
        rows.append({
            "file": file_path.name,
            "product": code,
            "target": str(target.relative_to(VS_ROOT)).replace(os.sep, "/"),
            "url": public_url_for(target),
            "status": action,
        })
    return rows


def resolve_upload_path(relative_path: str) -> Path:
    clean = relative_path.strip().strip("/\\")
    if not clean:
        raise ValueError("Enter a destination path after VS/.")
    parts = [part for part in re.split(r"[\\/]+", clean) if part and part not in {".", ".."}]
    if not parts:
        raise ValueError("Enter a valid destination path.")
    safe_parts = []
    for part in parts:
        safe = "".join(char for char in part if char.isalnum() or char in {" ", "_", "-", "."}).strip()
        if not safe:
            raise ValueError("Destination path contains an invalid folder name.")
        safe_parts.append(safe)
    destination = VS_ROOT.joinpath(*safe_parts).resolve()
    if not is_under_root(destination):
        raise ValueError("Destination must stay inside the VS folder.")
    return destination


def safe_upload_filename(filename: str) -> str:
    name = Path(filename or "").name
    safe = "".join(char for char in name if char.isalnum() or char in {" ", "_", "-", "."}).strip()
    if not safe or "." not in safe:
        raise ValueError("Invalid upload filename.")
    return safe


def detect_years_in_image(path: Path) -> dict[str, object]:
    if path.suffix.lower() not in IMAGE_EXTENSIONS:
        return {"checked": False, "years": [], "message": "Year check only runs on image files."}
    tesseract = shutil.which("tesseract")
    if not tesseract:
        return {"checked": False, "years": [], "message": "Tesseract OCR is not installed on the server."}
    if Image is None:
        return {"checked": False, "years": [], "message": "Pillow is not installed on the server."}

    variants: list[Path] = []
    try:
        with Image.open(path) as source:
            image = source.convert("RGB")
            max_width = 1800
            if image.width < max_width:
                scale = max_width / image.width
                image = image.resize((max_width, int(image.height * scale)), Image.Resampling.LANCZOS)
            gray = ImageOps.grayscale(image)
            contrast = ImageOps.autocontrast(gray)
            threshold = contrast.point(lambda p: 255 if p > 155 else 0)

            for suffix, candidate in (("gray", contrast), ("threshold", threshold)):
                temp = tempfile.NamedTemporaryFile(delete=False, suffix=f"-year-{suffix}.png")
                temp.close()
                candidate.save(temp.name)
                variants.append(Path(temp.name))

        texts = []
        for variant in variants:
            result = subprocess.run(
                [tesseract, str(variant), "stdout", "--psm", "6"],
                check=False,
                capture_output=True,
                text=True,
                timeout=15,
            )
            if result.stdout:
                texts.append(result.stdout)
        text = "\n".join(texts)
        years = sorted(set(YEAR_RE.findall(text)))
        snippet = " ".join(text.split())
        if len(snippet) > 220:
            snippet = snippet[:217] + "..."
        return {"checked": True, "years": years, "text": snippet, "message": f"Detected year(s): {', '.join(years)}" if years else "No year detected."}
    except subprocess.TimeoutExpired:
        return {"checked": False, "years": [], "message": "Year OCR timed out."}
    except Exception as exc:
        return {"checked": False, "years": [], "message": f"Year OCR failed: {exc}"}
    finally:
        for variant in variants:
            try:
                variant.unlink(missing_ok=True)
            except Exception:
                pass


def normalize_logo_text(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def iter_riegel_reference_files() -> list[Path]:
    files: list[Path] = []
    configured = os.environ.get("RIEGEL_REFERENCE_PATHS", "")
    paths = [Path(item.strip().strip('"')) for item in configured.split(";") if item.strip()]
    paths.extend(RIEGEL_REFERENCE_PATHS)
    for path in paths:
        try:
            if path.is_file() and path.suffix.lower() in LOGO_SCAN_EXTENSIONS:
                files.append(path)
            elif path.is_dir():
                files.extend(local_logo_scan_files(str(path), recursive=True)[:40])
        except Exception:
            continue
    cache_dir = Path(tempfile.gettempdir()) / "primex-riegel-references"
    for index, url in enumerate(RIEGEL_REFERENCE_URLS, start=1):
        try:
            cache_dir.mkdir(parents=True, exist_ok=True)
            parsed = urlparse(url)
            suffix = Path(parsed.path).suffix.lower() or ".jpg"
            target = cache_dir / f"reference-{index}{suffix}"
            if not target.exists() or target.stat().st_size < 1000:
                request = Request(url, headers={"User-Agent": "VS-Tools-Riegel-Logo-Scan/1.0"})
                with urlopen(request, timeout=20) as response:
                    target.write_bytes(response.read(5 * 1024 * 1024))
            if target.exists() and target.suffix.lower() in LOGO_SCAN_EXTENSIONS:
                files.append(target)
        except Exception:
            continue
    unique: dict[str, Path] = {}
    for path in files:
        unique[str(path.resolve()).lower()] = path
    return list(unique.values())[:60]


def image_for_template(path: Path, width: int = 240):
    with Image.open(path) as source:
        img = ImageOps.grayscale(source.convert("RGB"))
        img = ImageOps.autocontrast(img)
        if img.width > width:
            scale = width / img.width
            img = img.resize((width, max(1, int(img.height * scale))), Image.Resampling.LANCZOS)
        elif img.width < 120:
            scale = 120 / img.width
            img = img.resize((120, max(1, int(img.height * scale))), Image.Resampling.LANCZOS)
        return img


def template_crops(img) -> list[object]:
    w, h = img.size
    boxes = [
        (0, 0, w, h),
        (int(w * 0.2), int(h * 0.15), int(w * 0.8), int(h * 0.75)),
        (int(w * 0.25), int(h * 0.2), int(w * 0.75), int(h * 0.62)),
        (int(w * 0.28), int(h * 0.28), int(w * 0.72), int(h * 0.58)),
    ]
    crops = []
    for box in boxes:
        crop = img.crop(box)
        if crop.width >= 24 and crop.height >= 14:
            crops.append(crop)
    return crops


def normalized_template_score(candidate_img, template_img) -> float:
    if np is None:
        return normalized_template_score_pillow(candidate_img, template_img)
    cand = np.asarray(candidate_img, dtype=np.float32)
    tmpl = np.asarray(template_img, dtype=np.float32)
    th, tw = tmpl.shape
    ch, cw = cand.shape
    if th > ch or tw > cw or th < 8 or tw < 8:
        return 0.0
    tmpl = tmpl - tmpl.mean()
    tmpl_norm = float(np.sqrt((tmpl * tmpl).sum()))
    if tmpl_norm < 1:
        return 0.0
    best = 0.0
    step = max(4, min(th, tw) // 5)
    for y in range(0, ch - th + 1, step):
        for x in range(0, cw - tw + 1, step):
            patch = cand[y:y + th, x:x + tw]
            patch = patch - patch.mean()
            patch_norm = float(np.sqrt((patch * patch).sum()))
            if patch_norm < 1:
                continue
            score = float((patch * tmpl).sum() / (patch_norm * tmpl_norm))
            if score > best:
                best = score
    return best


def normalized_template_score_pillow(candidate_img, template_img) -> float:
    candidate_img = candidate_img.convert("L")
    template_img = template_img.convert("L")
    tw, th = template_img.size
    cw, ch = candidate_img.size
    if th > ch or tw > cw or th < 8 or tw < 8:
        return 0.0
    tmpl_values = [float(v) for v in template_img.getdata()]
    tmpl_mean = sum(tmpl_values) / len(tmpl_values)
    tmpl_centered = [v - tmpl_mean for v in tmpl_values]
    tmpl_norm = math.sqrt(sum(v * v for v in tmpl_centered))
    if tmpl_norm < 1:
        return 0.0
    best = 0.0
    step = max(4, min(th, tw) // 4)
    for y in range(0, ch - th + 1, step):
        for x in range(0, cw - tw + 1, step):
            patch = candidate_img.crop((x, y, x + tw, y + th))
            values = [float(v) for v in patch.getdata()]
            mean = sum(values) / len(values)
            centered = [v - mean for v in values]
            norm = math.sqrt(sum(v * v for v in centered))
            if norm < 1:
                continue
            score = sum(a * b for a, b in zip(centered, tmpl_centered)) / (norm * tmpl_norm)
            if score > best:
                best = score
    return best


def detect_riegel_reference_match(path: Path) -> dict[str, object]:
    if Image is None:
        return {"checked": False, "match": False, "score": 0.0, "message": "Visual reference scan needs Pillow."}
    references = iter_riegel_reference_files()
    if not references:
        return {"checked": False, "match": False, "score": 0.0, "message": "No Riegel reference image found."}
    try:
        candidate = image_for_template(path, width=320)
        best_score = 0.0
        best_ref = ""
        for ref in references:
            try:
                reference = image_for_template(ref, width=240)
                for crop in template_crops(reference):
                    for scale in (0.75, 1.0, 1.25):
                        size = (max(8, int(crop.width * scale)), max(8, int(crop.height * scale)))
                        template = crop.resize(size, Image.Resampling.LANCZOS)
                        score = normalized_template_score(candidate, template)
                        if score > best_score:
                            best_score = score
                            best_ref = str(ref)
            except Exception:
                continue
        match = best_score >= 0.9
        return {
            "checked": True,
            "match": match,
            "score": round(best_score, 3),
            "reference": best_ref,
            "message": f"Visual reference match score {best_score:.3f}." if match else f"Best visual reference score {best_score:.3f}.",
        }
    except Exception as exc:
        return {"checked": False, "match": False, "score": 0.0, "message": f"Visual reference scan failed: {exc}"}


def merge_riegel_scan_results(ocr_result: dict[str, object], visual_result: dict[str, object]) -> dict[str, object]:
    if visual_result.get("checked"):
        merged = dict(visual_result)
        if ocr_result.get("text"):
            merged["text"] = ocr_result.get("text", "")
        if visual_result.get("match"):
            merged["message"] = f"Visual reference matched after OCR missed the text. {visual_result.get('message', '')}"
        else:
            merged["message"] = f"{ocr_result.get('message', '')} {visual_result.get('message', '')}".strip()
        return merged
    return ocr_result


def detect_riegel_logo_in_image(path: Path) -> dict[str, object]:
    if path.suffix.lower() not in LOGO_SCAN_EXTENSIONS:
        return {"checked": False, "match": False, "message": "Logo scan only runs on image files."}
    tesseract = shutil.which("tesseract")
    if Image is None:
        return {"checked": False, "match": False, "message": "Pillow is not installed on the server."}
    if not tesseract:
        return detect_riegel_reference_match(path)

    variants: list[Path] = []
    try:
        with Image.open(path) as source:
            image = source.convert("RGB")
            max_width = 2200
            if image.width < max_width:
                scale = max_width / image.width
                image = image.resize((max_width, int(image.height * scale)), Image.Resampling.LANCZOS)
            gray = ImageOps.grayscale(image)
            contrast = ImageOps.autocontrast(gray)
            threshold_dark = contrast.point(lambda p: 255 if p > 145 else 0)
            threshold_light = contrast.point(lambda p: 0 if p > 185 else 255)
            variants_to_save = [
                ("gray", contrast),
                ("dark", threshold_dark),
                ("light", threshold_light),
            ]
            for suffix, candidate in variants_to_save:
                temp = tempfile.NamedTemporaryFile(delete=False, suffix=f"-riegel-{suffix}.png")
                temp.close()
                candidate.save(temp.name)
                variants.append(Path(temp.name))

        texts = []
        for variant in variants:
            for psm in ("6", "11"):
                result = subprocess.run(
                    [tesseract, str(variant), "stdout", "--psm", psm],
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=15,
                )
                if result.stdout:
                    texts.append(result.stdout)
        text = "\n".join(texts)
        normalized = normalize_logo_text(text)
        match = "riegel" in normalized or bool(RIEGEL_RE.search(text))
        snippet = " ".join(text.split())
        if len(snippet) > 260:
            snippet = snippet[:257] + "..."
        return {
            "checked": True,
            "match": match,
            "text": snippet,
            "message": "Riegel logo text detected." if match else "No Riegel logo text detected.",
        } if match else merge_riegel_scan_results(
            {"checked": True, "match": False, "text": snippet, "message": "No Riegel logo text detected."},
            detect_riegel_reference_match(path),
        )
    except subprocess.TimeoutExpired:
        return {"checked": False, "match": False, "message": "Riegel OCR timed out."}
    except Exception as exc:
        fallback = detect_riegel_reference_match(path)
        if fallback.get("checked"):
            fallback["message"] = f"OCR failed, used visual scan. {fallback.get('message', '')}"
            return fallback
        return {"checked": False, "match": False, "message": f"Riegel OCR failed: {exc}"}
    finally:
        for variant in variants:
            try:
                variant.unlink(missing_ok=True)
            except Exception:
                pass


def local_logo_scan_files(folder_text: str, recursive: bool) -> list[Path]:
    folder = Path(folder_text.strip().strip('"'))
    if not folder.exists() or not folder.is_dir():
        raise ValueError("Folder does not exist or is not a folder.")
    files = folder.rglob("*") if recursive else folder.iterdir()
    return sorted([path for path in files if path.is_file() and path.suffix.lower() in LOGO_SCAN_EXTENSIONS], key=lambda path: str(path).lower())


def download_scan_image(url: str) -> Path:
    parsed = urlparse(url)
    suffix = Path(parsed.path).suffix.lower()
    if suffix not in LOGO_SCAN_EXTENSIONS:
        suffix = ".jpg"
    request = Request(url, headers={"User-Agent": "VS-Tools-Riegel-Logo-Scan/1.0"})
    with urlopen(request, timeout=20) as response:
        data = response.read(15 * 1024 * 1024 + 1)
    if len(data) > 15 * 1024 * 1024:
        raise ValueError("Image is larger than 15 MB.")
    temp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    try:
        temp.write(data)
        temp.close()
        return Path(temp.name)
    except Exception:
        temp.close()
        Path(temp.name).unlink(missing_ok=True)
        raise


def scan_riegel_sources(url_text: str, folder_text: str, recursive: bool) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    urls = [line.strip() for line in re.split(r"[\r\n]+", url_text or "") if line.strip()]
    seen_urls: set[str] = set()
    for url in urls[:1000]:
        if url in seen_urls:
            continue
        seen_urls.add(url)
        temp_path: Path | None = None
        try:
            if not url.lower().startswith(("http://", "https://")):
                raise ValueError("Only http/https image URLs are supported.")
            temp_path = download_scan_image(url)
            result = detect_riegel_logo_in_image(temp_path)
            rows.append({
                "source": "url",
                "file": Path(urlparse(url).path).name or url,
                "path": url,
                "url": url,
                "status": "riegel found" if result.get("match") else ("no riegel" if result.get("checked") else "skipped"),
                "match": bool(result.get("match")),
                "message": result.get("message", ""),
                "text": result.get("text", ""),
                "score": result.get("score", ""),
                "reference": result.get("reference", ""),
            })
        except Exception as exc:
            rows.append({"source": "url", "file": Path(urlparse(url).path).name or url, "path": url, "url": url, "status": "error", "match": False, "message": str(exc)})
        finally:
            if temp_path:
                try:
                    temp_path.unlink(missing_ok=True)
                except Exception:
                    pass

    if folder_text.strip():
        files = local_logo_scan_files(folder_text, recursive)
        for path in files[:2000]:
            try:
                result = detect_riegel_logo_in_image(path)
                rows.append({
                    "source": "folder",
                    "file": path.name,
                    "path": str(path),
                    "status": "riegel found" if result.get("match") else ("no riegel" if result.get("checked") else "skipped"),
                    "match": bool(result.get("match")),
                    "message": result.get("message", ""),
                    "text": result.get("text", ""),
                    "score": result.get("score", ""),
                    "reference": result.get("reference", ""),
                })
            except Exception as exc:
                rows.append({"source": "folder", "file": path.name, "path": str(path), "status": "error", "match": False, "message": str(exc)})
    if not rows:
        raise ValueError("Paste image URLs or enter a folder path to scan.")
    return rows


def riegel_url_scan_row(url: str) -> dict[str, object]:
    temp_path: Path | None = None
    try:
        if not url.lower().startswith(("http://", "https://")):
            raise ValueError("Only http/https image URLs are supported.")
        temp_path = download_scan_image(url)
        result = detect_riegel_logo_in_image(temp_path)
        return {
            "source": "url",
            "file": Path(urlparse(url).path).name or url,
            "path": url,
            "url": url,
            "status": "riegel found" if result.get("match") else ("no riegel" if result.get("checked") else "skipped"),
            "match": bool(result.get("match")),
            "message": result.get("message", ""),
            "text": result.get("text", ""),
            "score": result.get("score", ""),
            "reference": result.get("reference", ""),
        }
    except Exception as exc:
        return {"source": "url", "file": Path(urlparse(url).path).name or url, "path": url, "url": url, "status": "error", "match": False, "message": str(exc)}
    finally:
        if temp_path:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass


def riegel_file_scan_row(path: Path) -> dict[str, object]:
    try:
        result = detect_riegel_logo_in_image(path)
        return {
            "source": "folder",
            "file": path.name,
            "path": str(path),
            "status": "riegel found" if result.get("match") else ("no riegel" if result.get("checked") else "skipped"),
            "match": bool(result.get("match")),
            "message": result.get("message", ""),
            "text": result.get("text", ""),
            "score": result.get("score", ""),
            "reference": result.get("reference", ""),
        }
    except Exception as exc:
        return {"source": "folder", "file": path.name, "path": str(path), "status": "error", "match": False, "message": str(exc)}


def update_riegel_job_progress(job_id: str, rows: list[dict[str, object]], index: int, total: int, current: str) -> None:
    found = sum(1 for row in rows if row.get("match") is True)
    issues = sum(1 for row in rows if row.get("status") in {"error", "skipped"})
    percent = 8 + int((index / max(1, total)) * 88)
    update_job(
        job_id,
        phase="scanning",
        percent=min(96, percent),
        message=f"Scanning {index}/{total}: {current}",
        rows=list(rows),
        scanned=len(rows),
        found=found,
        issues=issues,
    )


def run_riegel_scan_job(job_id: str, url_text: str, folder_text: str, recursive: bool) -> None:
    try:
        urls = [line.strip() for line in re.split(r"[\r\n]+", url_text or "") if line.strip()]
        unique_urls = []
        seen_urls: set[str] = set()
        for url in urls[:1000]:
            if url not in seen_urls:
                unique_urls.append(url)
                seen_urls.add(url)
        files = local_logo_scan_files(folder_text, recursive)[:2000] if folder_text.strip() else []
        total = len(unique_urls) + len(files)
        if total == 0:
            raise ValueError("Paste image URLs or enter a folder path to scan.")
        rows: list[dict[str, object]] = []
        update_job(job_id, phase="scanning", percent=8, message=f"Preparing to scan {total} image(s)...", rows=[], scanned=0, found=0, issues=0)
        index = 0
        for url in unique_urls:
            index += 1
            name = Path(urlparse(url).path).name or url
            update_riegel_job_progress(job_id, rows, index, total, name)
            rows.append(riegel_url_scan_row(url))
            update_riegel_job_progress(job_id, rows, index, total, name)
        for path in files:
            index += 1
            update_riegel_job_progress(job_id, rows, index, total, path.name)
            rows.append(riegel_file_scan_row(path))
            update_riegel_job_progress(job_id, rows, index, total, path.name)
        found = sum(1 for row in rows if row.get("match") is True)
        issues = sum(1 for row in rows if row.get("status") in {"error", "skipped"})
        update_job(
            job_id,
            done=True,
            ok=True,
            phase="complete",
            percent=100,
            message=f"Finished. Riegel found in {found} image(s).",
            rows=rows,
            scanned=len(rows),
            found=found,
            issues=issues,
        )
    except Exception as exc:
        update_job(
            job_id,
            done=True,
            ok=False,
            phase="error",
            percent=100,
            message=str(exc),
            issues=1,
        )


def parse_sku_for_excel(sku: str, prefix: str, strip_suffix: str, top_folder: str = "") -> tuple[str, str]:
    value = str(sku or "").strip()
    quantity = "1"
    match = re.match(r"^(01|03|06|12)-(.+)$", value, re.IGNORECASE)
    if match:
        quantity = str(int(match.group(1)))
        value = match.group(2)
    clean = "".join(char for char in value if char.isalnum())
    if strip_suffix and clean.lower().endswith(strip_suffix.lower()):
        clean = clean[: -len(strip_suffix)]
    if prefix and clean.upper().startswith(prefix.upper()):
        clean = clean[len(prefix):]
    elif top_folder and clean.upper().startswith(top_folder.upper()):
        clean = clean[len(top_folder):]
    if not clean:
        raise ValueError(f"Could not detect product folder from SKU: {sku}")
    return clean, quantity


def column_letters_to_number(letters: str) -> int:
    number = 0
    for char in letters.upper():
        number = number * 26 + ord(char) - 64
    return number


def column_number_to_letters(number: int) -> str:
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def split_cell_ref(ref: str) -> tuple[str, int]:
    match = re.match(r"^([A-Z]+)(\d+)$", ref.upper())
    if not match:
        raise ValueError(f"Invalid cell reference: {ref}")
    return match.group(1), int(match.group(2))


def fill_excel_urls(
    source_file: Path,
    original_name: str,
    top_folder: str,
    prefix: str,
    strip_suffix: str,
    cache_bust: bool = False,
) -> tuple[Path, dict[str, int]]:
    if not is_safe_name(top_folder):
        raise ValueError("Select a valid Hetzner top folder.")
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    ET.register_namespace("", ns_main)
    ET.register_namespace("r", ns_rel)
    with zipfile.ZipFile(source_file, "r") as zin:
        workbook_xml = ET.fromstring(zin.read("xl/workbook.xml"))
        rels_xml = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        sheet_rel_id = None
        for sheet in workbook_xml.findall(f".//{{{ns_main}}}sheet"):
            if sheet.attrib.get("name") == "Vorlage":
                sheet_rel_id = sheet.attrib.get(f"{{{ns_rel}}}id")
                break
        if not sheet_rel_id:
            raise ValueError("Sheet 'Vorlage' was not found in the workbook.")
        sheet_target = None
        for rel in rels_xml.findall(f"{{{ns_pkg_rel}}}Relationship"):
            if rel.attrib.get("Id") == sheet_rel_id:
                sheet_target = rel.attrib.get("Target")
                break
        if not sheet_target:
            raise ValueError("Could not locate the 'Vorlage' worksheet file.")
        sheet_path = "xl/" + sheet_target.lstrip("/")
        shared_path = "xl/sharedStrings.xml"
        if shared_path not in zin.namelist():
            raise ValueError("Workbook does not contain shared strings.")
        shared_xml = ET.fromstring(zin.read(shared_path))
        sheet_xml = ET.fromstring(zin.read(sheet_path))

        def shared_text(si) -> str:
            return "".join(node.text or "" for node in si.iter(f"{{{ns_main}}}t"))

        shared_strings = [shared_text(si) for si in shared_xml.findall(f"{{{ns_main}}}si")]
        shared_index = {value: idx for idx, value in enumerate(shared_strings)}

        def cell_text(cell) -> str:
            cell_type = cell.attrib.get("t")
            if cell_type == "s":
                value = cell.find(f"{{{ns_main}}}v")
                return shared_strings[int(value.text)] if value is not None and value.text else ""
            if cell_type == "inlineStr":
                return "".join(node.text or "" for node in cell.iter(f"{{{ns_main}}}t"))
            value = cell.find(f"{{{ns_main}}}v")
            return value.text if value is not None and value.text else ""

        def add_shared_string(value: str) -> int:
            if value in shared_index:
                return shared_index[value]
            si = ET.SubElement(shared_xml, f"{{{ns_main}}}si")
            text = ET.SubElement(si, f"{{{ns_main}}}t")
            text.text = value
            shared_index[value] = len(shared_strings)
            shared_strings.append(value)
            return shared_index[value]

        sheet_data = sheet_xml.find(f"{{{ns_main}}}sheetData")
        if sheet_data is None:
            raise ValueError("Could not read worksheet rows.")

        rows = {}
        for row in sheet_data.findall(f"{{{ns_main}}}row"):
            row_num = int(row.attrib.get("r", "0"))
            rows[row_num] = row

        def cells_by_ref(row):
            return {cell.attrib.get("r"): cell for cell in row.findall(f"{{{ns_main}}}c")}

        header_row = rows.get(4)
        if header_row is None:
            raise ValueError("Could not find header row 4.")
        main_col = None
        other_col = None
        for cell in header_row.findall(f"{{{ns_main}}}c"):
            ref = cell.attrib.get("r", "")
            text = cell_text(cell)
            if text == "URL des Hauptbildes" and main_col is None:
                main_col = column_letters_to_number(split_cell_ref(ref)[0])
            if text == "Andere Bild-URL" and other_col is None:
                other_col = column_letters_to_number(split_cell_ref(ref)[0])
        if not main_col or not other_col:
            raise ValueError("Could not find image URL columns in row 4.")

        def set_text_cell(row, column_number: int, value: str) -> None:
            ref = f"{column_number_to_letters(column_number)}{row.attrib['r']}"
            existing = cells_by_ref(row).get(ref)
            if existing is None:
                existing = ET.Element(f"{{{ns_main}}}c", {"r": ref})
                row.append(existing)
            style = existing.attrib.get("s")
            existing.attrib.clear()
            existing.attrib["r"] = ref
            if style:
                existing.attrib["s"] = style
            existing.attrib["t"] = "s"
            for child in list(existing):
                existing.remove(child)
            value_node = ET.SubElement(existing, f"{{{ns_main}}}v")
            value_node.text = str(add_shared_string(value))

        max_row = max(rows) if rows else 0
        changed = 0
        missing = 0
        skipped = 0
        for row_num in range(7, max_row + 1):
            row = rows.get(row_num)
            if row is None:
                continue
            sku_cell = cells_by_ref(row).get(f"A{row_num}")
            sku = cell_text(sku_cell) if sku_cell is not None else ""
            if not sku:
                continue
            try:
                folder, quantity = parse_sku_for_excel(str(sku), prefix, strip_suffix, top_folder)
            except ValueError:
                skipped += 1
                continue
            product_folder = VS_ROOT / top_folder / folder
            main_file = product_folder / f"{quantity}.jpg"
            one_file = product_folder / "1.jpg"
            if not main_file.exists() or not one_file.exists():
                missing += 1
            main_url = public_url_for(main_file, cache_bust)
            one_url = public_url_for(one_file, cache_bust)
            set_text_cell(row, main_col, main_url)
            set_text_cell(row, other_col, one_url)
            changed += 1

        shared_xml.attrib["count"] = str(max(int(shared_xml.attrib.get("count", "0")), len(shared_strings)))
        shared_xml.attrib["uniqueCount"] = str(len(shared_strings))

        suffix = Path(original_name or "amazon-template.xlsm").suffix or ".xlsm"
        stem = Path(original_name or "amazon-template").stem
        safe_stem = "".join(char for char in stem if char.isalnum() or char in {" ", "_", "-"}).strip() or "amazon-template"
        output = Path(tempfile.gettempdir()) / f"{safe_stem}-with-urls{suffix}"
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == sheet_path:
                    zout.writestr(item, ET.tostring(sheet_xml, encoding="utf-8", xml_declaration=True))
                elif item.filename == shared_path:
                    zout.writestr(item, ET.tostring(shared_xml, encoding="utf-8", xml_declaration=True))
                else:
                    zout.writestr(item, zin.read(item.filename))
        return output, {"changed": changed, "missing": missing, "skipped": skipped}


GEWICHT_QUANTITIES = (1, 3, 6, 12)
GEWICHT_CARTONS = (180, 540, 1120, 2120)


def parse_gewicht_value(value: object, excel_row: int, header_unit: str = "") -> tuple[Decimal, str]:
    if value is None or isinstance(value, bool):
        raise ValueError(f"Row {excel_row}: Gewicht is empty or invalid.")
    raw = str(value).strip().replace(" ", "").replace("\u00a0", "")
    match = re.fullmatch(r"([+-]?(?:\d+[.,]?)*\d+)(kg|g|mg)?", raw, re.IGNORECASE)
    if not match:
        raise ValueError(f"Row {excel_row}: invalid Gewicht value {value!r}.")
    number_text, explicit_unit = match.groups()
    if "," in number_text and "." in number_text:
        number_text = number_text.replace(".", "").replace(",", ".")
    else:
        number_text = number_text.replace(",", ".")
    try:
        number = Decimal(number_text)
    except InvalidOperation as exc:
        raise ValueError(f"Row {excel_row}: invalid Gewicht value {value!r}.") from exc
    if not number.is_finite() or number < 0:
        raise ValueError(f"Row {excel_row}: Gewicht must be a non-negative number.")
    unit = (explicit_unit or header_unit).lower()
    if not unit:
        whole = number_text.lstrip("+-").split(".", 1)[0]
        unit = "mg" if len(whole) >= 2 else "g"
    if unit == "kg":
        return number * 1000000, unit
    if unit == "g":
        return number * 1000, unit
    return number, unit


def gewicht_excel_number(number: Decimal) -> int | float:
    return int(number) if number == number.to_integral_value() else float(number)


def build_gewicht_workbook(source_file: Path, prefix: str = "AM", input_unit: str = "g", code_suffix: str = "") -> tuple[Path, int]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise ValueError("Excel support is not installed on the server.") from exc
    prefix = prefix.strip()
    if not prefix or not re.fullmatch(r"[A-Za-z0-9]+", prefix):
        raise ValueError("Prefix must contain only letters and numbers.")
    input_unit = input_unit.strip().lower()
    if input_unit not in {"g", "mg"}:
        raise ValueError("Choose g or mg for the uploaded weight.")
    code_suffix = code_suffix.strip()
    if code_suffix and not re.fullmatch(r"[A-Za-z0-9]+", code_suffix):
        raise ValueError("The optional code suffix must contain only letters and numbers.")
    source = load_workbook(source_file, read_only=True, data_only=True)
    try:
        sheet = source.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError("Excel file is empty.") from exc
        first = str(header[0] or "").strip().casefold() if len(header) > 0 else ""
        second = str(header[1] or "").strip().casefold() if len(header) > 1 else ""
        if first not in {"artikel nr", "artikel nr.", "artikelnummer", "artikel-nr", "artikel-nr."} or not second.startswith("gewicht"):
            raise ValueError("First row must contain 'Artikel Nr' and 'Gewicht' in columns A and B.")
        has_total_columns = (len(header) > 5 and str(header[4] or "").strip().casefold().startswith("karton")
                             and str(header[5] or "").strip().casefold().startswith("gewicht"))
        has_product_column = len(header) > 3 and str(header[3] or "").strip().casefold().startswith("produktgewicht")
        input_rows = []
        for excel_row, row in enumerate(rows, start=2):
            article = row[0] if len(row) > 0 else None
            weight = row[1] if len(row) > 1 else None
            if article is None and weight is None:
                continue
            if article is None or str(article).strip() == "":
                raise ValueError(f"Row {excel_row}: Artikel Nr is missing.")
            if weight is None or str(weight).strip() == "":
                milligrams, unit = None, ""
                quantity_value = row[2] if len(row) > 2 else None
                try:
                    quantity = Decimal(str(quantity_value).strip())
                    if not quantity.is_finite() or quantity <= 0:
                        raise InvalidOperation
                    if has_total_columns and len(row) > 5 and row[4] is not None and row[5] is not None:
                        total, _ = parse_gewicht_value(row[5], excel_row, "mg")
                        carton, _ = parse_gewicht_value(row[4], excel_row, "mg")
                        milligrams = (total - carton) / quantity
                    elif has_product_column and len(row) > 3 and row[3] is not None:
                        product, _ = parse_gewicht_value(row[3], excel_row, "mg")
                        milligrams = product / quantity
                    if milligrams is not None and milligrams < 0:
                        raise ValueError(f"Row {excel_row}: derived Gewicht must not be negative.")
                    if milligrams is not None:
                        weight, unit = gewicht_excel_number(milligrams), "mg"
                except (InvalidOperation, TypeError):
                    pass
            else:
                milligrams, unit = parse_gewicht_value(weight, excel_row, input_unit)
            input_rows.append((article, weight, milligrams, unit))
        if not input_rows:
            raise ValueError("No product rows found in the Excel file.")
        source_units = {unit for _, _, _, unit in input_rows if unit}
        has_conversion = any(unit != "mg" for unit in source_units)
        output = Workbook()
        result = output.active
        result.title = "GEWICHT"
        original_unit = next(iter(source_units)) if len(source_units) == 1 else ""
        original_header = f"Gewicht ({original_unit})" if original_unit else ("Gewicht (mg)" if not source_units else "Gewicht (siç u dërgua)")
        headers = ["Artikel Nr", "Artikel Code", original_header]
        if has_conversion:
            headers.append("Gewicht (mg)")
        headers.extend(["Anzahl Flaschen", "Produktgewicht (mg)", "Karton (mg)", "GEWICHT (mg)"])
        result.append(headers)
        for cell in result[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
        for count, (article, weight, milligrams, unit) in enumerate(input_rows):
            variant = count % 4
            quantity = GEWICHT_QUANTITIES[variant]
            carton = GEWICHT_CARTONS[variant]
            product = milligrams * quantity if milligrams is not None else None
            total = product + carton if product is not None else None
            article_text = str(int(article)) if isinstance(article, float) and article.is_integer() else str(article).strip()
            values = [article, f"{quantity:02d}-{prefix}{article_text}{code_suffix}", weight]
            if has_conversion:
                values.append(gewicht_excel_number(milligrams) if unit != "mg" and milligrams is not None else None)
            values.extend([quantity, gewicht_excel_number(product) if product is not None else None,
                           carton if milligrams is not None else None,
                           gewicht_excel_number(total) if total is not None else None])
            result.append(values)
        for column in result.columns:
            result.column_dimensions[column[0].column_letter].width = max(18, min(28, len(str(column[0].value)) + 3))
        result.freeze_panes = "A2"
        result.auto_filter.ref = result.dimensions
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp:
            output_path = Path(temp.name)
        output.save(output_path)
        return output_path, len(input_rows)
    finally:
        source.close()


def fill_excel_urls_workbook(
    source_file: Path,
    original_name: str,
    top_folder: str,
    prefix: str,
    strip_suffix: str,
    cache_bust: bool = False,
) -> tuple[Path, dict[str, int]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise ValueError("Excel support is not installed on the server.") from exc
    if not is_safe_name(top_folder):
        raise ValueError("Select a valid Hetzner top folder.")
    workbook = load_workbook(source_file, keep_vba=True)
    if "Vorlage" not in workbook.sheetnames:
        raise ValueError("Sheet 'Vorlage' was not found in the workbook.")
    sheet = workbook["Vorlage"]
    main_col = None
    other_col = None
    for cell in sheet[4]:
        if cell.value == "URL des Hauptbildes" and main_col is None:
            main_col = cell.column
        if cell.value == "Andere Bild-URL" and other_col is None:
            other_col = cell.column
    if not main_col or not other_col:
        raise ValueError("Could not find image URL columns in row 4.")
    changed = 0
    missing = 0
    skipped = 0
    for row in range(7, sheet.max_row + 1):
        sku = sheet.cell(row=row, column=1).value
        if not sku:
            continue
        try:
            folder, quantity = parse_sku_for_excel(str(sku), prefix, strip_suffix, top_folder)
        except ValueError:
            skipped += 1
            continue
        product_folder = VS_ROOT / top_folder / folder
        main_file = product_folder / f"{quantity}.jpg"
        one_file = product_folder / "1.jpg"
        if not main_file.exists() or not one_file.exists():
            missing += 1
        main_url = public_url_for(main_file, cache_bust)
        one_url = public_url_for(one_file, cache_bust)
        sheet.cell(row=row, column=main_col).value = main_url
        sheet.cell(row=row, column=other_col).value = one_url
        changed += 1
    suffix = Path(original_name or "amazon-template.xlsm").suffix or ".xlsm"
    stem = Path(original_name or "amazon-template").stem
    safe_stem = "".join(char for char in stem if char.isalnum() or char in {" ", "_", "-"}).strip() or "amazon-template"
    output = Path(tempfile.gettempdir()) / f"{safe_stem}-with-urls{suffix}"
    workbook.save(output)
    return output, {"changed": changed, "missing": missing, "skipped": skipped}


def fill_excel_urls_zip_patch(
    source_file: Path,
    original_name: str,
    top_folder: str,
    prefix: str,
    strip_suffix: str,
    cache_bust: bool = False,
) -> tuple[Path, dict[str, int]]:
    if not is_safe_name(top_folder):
        raise ValueError("Select a valid Hetzner top folder.")
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(source_file, "r") as zin:
        workbook_xml = ET.fromstring(zin.read("xl/workbook.xml"))
        rels_xml = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        sheet_rel_id = None
        for sheet in workbook_xml.findall(f".//{{{ns_main}}}sheet"):
            if sheet.attrib.get("name") == "Vorlage":
                sheet_rel_id = sheet.attrib.get(f"{{{ns_rel}}}id")
                break
        if not sheet_rel_id:
            raise ValueError("Sheet 'Vorlage' was not found in the workbook.")
        sheet_target = None
        for rel in rels_xml.findall(f"{{{ns_pkg_rel}}}Relationship"):
            if rel.attrib.get("Id") == sheet_rel_id:
                sheet_target = rel.attrib.get("Target")
                break
        if not sheet_target:
            raise ValueError("Could not locate the 'Vorlage' worksheet file.")
        sheet_path = "xl/" + sheet_target.lstrip("/")
        if sheet_path not in zin.namelist() and sheet_target.startswith("/xl/"):
            sheet_path = sheet_target.lstrip("/")
        shared_strings = []
        if "xl/sharedStrings.xml" in zin.namelist():
            shared_xml = ET.fromstring(zin.read("xl/sharedStrings.xml"))
            for si in shared_xml.findall(f"{{{ns_main}}}si"):
                shared_strings.append("".join(node.text or "" for node in si.iter(f"{{{ns_main}}}t")))
        sheet_bytes = zin.read(sheet_path)
        sheet_root = ET.fromstring(sheet_bytes)

        def read_cell_text(cell) -> str:
            if cell is None:
                return ""
            cell_type = cell.attrib.get("t")
            if cell_type == "s":
                value = cell.find(f"{{{ns_main}}}v")
                return shared_strings[int(value.text)] if value is not None and value.text else ""
            if cell_type == "inlineStr":
                return "".join(node.text or "" for node in cell.iter(f"{{{ns_main}}}t"))
            value = cell.find(f"{{{ns_main}}}v")
            return value.text if value is not None and value.text else ""

        main_col = None
        other_col = None
        for row in sheet_root.findall(f".//{{{ns_main}}}row"):
            if row.attrib.get("r") != "4":
                continue
            for cell in row.findall(f"{{{ns_main}}}c"):
                ref = cell.attrib.get("r", "")
                text = read_cell_text(cell)
                if text == "URL des Hauptbildes" and main_col is None:
                    main_col = column_letters_to_number(split_cell_ref(ref)[0])
                if text == "Andere Bild-URL" and other_col is None:
                    other_col = column_letters_to_number(split_cell_ref(ref)[0])
        if not main_col or not other_col:
            raise ValueError("Could not find image URL columns in row 4.")

        updates: dict[str, str] = {}
        url_status_cache: dict[str, tuple[bool, str]] = {}
        changed = 0
        missing = 0
        skipped = 0
        checked = 0
        valid = 0
        invalid = 0
        missing_products: set[str] = set()
        row_url_pairs: list[tuple[str, str]] = []
        pending_url_checks: set[str] = set()
        for row in sheet_root.findall(f".//{{{ns_main}}}row"):
            row_number = int(row.attrib.get("r", "0") or "0")
            if row_number < 7:
                continue
            sku_cell = None
            for cell in row.findall(f"{{{ns_main}}}c"):
                if cell.attrib.get("r") == f"A{row_number}":
                    sku_cell = cell
                    break
            sku = read_cell_text(sku_cell)
            if not sku:
                continue
            try:
                folder, quantity = parse_sku_for_excel(sku, prefix, strip_suffix, top_folder)
            except ValueError:
                skipped += 1
                continue
            product_folder = VS_ROOT / top_folder / folder
            main_file = product_folder / f"{quantity}.jpg"
            one_file = product_folder / "1.jpg"
            if not main_file.exists() or not one_file.exists():
                missing += 1
                missing_products.add(folder)
            main_url = public_url_for(main_file, cache_bust)
            one_url = public_url_for(one_file, cache_bust)
            if main_url not in url_status_cache:
                if main_file.exists():
                    pending_url_checks.add(main_url)
                else:
                    url_status_cache[main_url] = (False, "ERROR local file missing")
            if one_url not in url_status_cache:
                if one_file.exists():
                    pending_url_checks.add(one_url)
                else:
                    url_status_cache[one_url] = (False, "ERROR local file missing")
            row_url_pairs.append((main_url, one_url))
            updates[f"{column_number_to_letters(main_col)}{row_number}"] = main_url
            updates[f"{column_number_to_letters(other_col)}{row_number}"] = one_url
            changed += 1

        # URL checks are network-bound. Run them concurrently so large workbooks
        # do not exceed the reverse-proxy/Cloudflare request timeout.
        if pending_url_checks:
            with ThreadPoolExecutor(max_workers=16, thread_name_prefix="excel-url") as pool:
                futures = {pool.submit(check_public_url, url): url for url in pending_url_checks}
                for future in as_completed(futures):
                    url = futures[future]
                    try:
                        url_status_cache[url] = future.result()
                    except Exception as exc:
                        url_status_cache[url] = (False, f"ERROR URL check: {exc}")

        for main_url, one_url in row_url_pairs:
            main_ok, _ = url_status_cache[main_url]
            one_ok, _ = url_status_cache[one_url]
            checked += 2
            valid += int(main_ok) + int(one_ok)
            invalid += int(not main_ok) + int(not one_ok)

        def xml_escape(value: str) -> str:
            return html.escape(value, quote=False)

        def patch_cell(xml: str, ref: str, value: str) -> str:
            row_number = split_cell_ref(ref)[1]
            row_pattern = re.compile(rf"(<row\b[^>]*\br=\"{row_number}\"[^>]*>)(.*?)(</row>)", re.DOTALL)
            row_match = row_pattern.search(xml)
            if not row_match:
                return xml
            row_open, row_body, row_close = row_match.groups()
            cell_pattern = re.compile(rf"<c\b([^>]*\br=\"{re.escape(ref)}\"[^>]*)>(.*?)</c>|<c\b([^>]*\br=\"{re.escape(ref)}\"[^>]*)/>", re.DOTALL)
            cell_match = cell_pattern.search(row_body)
            attrs = cell_match.group(1) or cell_match.group(3) if cell_match else f' r="{ref}"'
            style_match = re.search(r'\bs="[^"]*"', attrs)
            style = f" {style_match.group(0)}" if style_match else ""
            new_cell = f'<c r="{ref}"{style} t="inlineStr"><is><t>{xml_escape(value)}</t></is></c>'
            if cell_match:
                row_body = row_body[:cell_match.start()] + new_cell + row_body[cell_match.end():]
            else:
                target_col = column_letters_to_number(split_cell_ref(ref)[0])
                insert_at = len(row_body)
                for existing in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', row_body):
                    if column_letters_to_number(existing.group(1)) > target_col:
                        insert_at = existing.start()
                        break
                row_body = row_body[:insert_at] + new_cell + row_body[insert_at:]
            return xml[:row_match.start()] + row_open + row_body + row_close + xml[row_match.end():]

        sheet_text = sheet_bytes.decode("utf-8")
        for ref, value in updates.items():
            sheet_text = patch_cell(sheet_text, ref, value)

        suffix = Path(original_name or "amazon-template.xlsm").suffix or ".xlsm"
        stem = Path(original_name or "amazon-template").stem
        safe_stem = "".join(char for char in stem if char.isalnum() or char in {" ", "_", "-"}).strip() or "amazon-template"
        output = Path(tempfile.gettempdir()) / f"{safe_stem}-with-urls{suffix}"
        with zipfile.ZipFile(output, "w") as zout:
            for item in zin.infolist():
                data = sheet_text.encode("utf-8") if item.filename == sheet_path else zin.read(item.filename)
                zout.writestr(item, data)
        return output, {
            "changed": changed,
            "missing": missing,
            "missing_products": sorted(missing_products),
            "skipped": skipped,
            "checked": checked,
            "valid": valid,
            "invalid": invalid,
        }


def upload_parts(relative_path: str) -> list[str]:
    destination = resolve_upload_path(relative_path)
    return list(destination.relative_to(VS_ROOT).parts)


STYLE = """
<style>
:root{--bg:#f5f7fb;--panel:#fff;--ink:#111827;--muted:#667085;--line:#d7dde8;--accent:#0f766e;--blue:#1d4ed8}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:Segoe UI,Arial,sans-serif;letter-spacing:0}
.app{display:grid;grid-template-columns:280px 1fr;min-height:100vh}.side{background:#111827;color:white;padding:22px 16px}.brand-wrap{display:grid;gap:12px;margin-bottom:42px}.logo-row{display:flex;align-items:center}.brand-logo{display:block;object-fit:contain}.brand-logo.primex{max-width:228px;max-height:64px}.brand{font-size:20px;font-weight:800}.meta{font-size:12px;color:#a7b0c0;word-break:break-word;margin-bottom:18px}
.nav a{display:flex;align-items:center;gap:10px;color:#e5e7eb;text-decoration:none;padding:11px 10px;border-radius:8px;margin:5px 0;font-weight:700}.nav a.active,.nav a:hover{background:#263244}.nav svg{width:18px;height:18px;stroke:currentColor;stroke-width:2;fill:none;stroke-linecap:round;stroke-linejoin:round;flex:0 0 auto}
main{padding:28px;max-width:1500px}.panel{background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:18px;margin-bottom:16px;box-shadow:0 10px 28px rgba(16,24,40,.07)}
h1{font-size:28px;margin:0 0 5px}.hint{color:var(--muted);font-size:14px;line-height:1.45}.grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px}.field{display:grid;gap:7px}.field.full{grid-column:1/-1}
label{font-size:13px;font-weight:750}select,input{width:100%;border:1px solid var(--line);border-radius:8px;padding:10px 11px;font:inherit;background:white}.checks{display:flex;gap:18px;align-items:center;flex-wrap:wrap}.check{display:flex;align-items:center;gap:8px}.check input{width:auto}
datalist option{font:inherit}
button{border:0;background:var(--accent);color:white;border-radius:8px;padding:11px 15px;font-weight:750;cursor:pointer;min-height:42px}button:disabled{opacity:.55;cursor:not-allowed}.secondary{background:var(--blue)}.mutedbtn{background:#475467}
.actions{display:flex;gap:10px;flex-wrap:wrap;margin-top:16px}.cards{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px}.card{border:1px solid var(--line);border-radius:8px;padding:12px;background:#f8fafc}.card .n{font-weight:800;font-size:22px}.card .t{color:var(--muted);font-size:12px}
.toolbar{display:flex;gap:12px;align-items:end;flex-wrap:wrap;margin:0 0 12px}.toolbar .field{min-width:210px}.thumb{width:72px;height:72px;object-fit:contain;background:#fff;border:1px solid var(--line);border-radius:6px}.thumb-cell{width:90px}.login-page{min-height:100vh;display:grid;place-items:center;padding:24px}.login-card{width:min(420px,100%);background:white;border:1px solid var(--line);border-radius:8px;padding:24px;box-shadow:0 18px 40px rgba(16,24,40,.12)}.login-logo{display:block;width:100%;height:96px;object-fit:contain;margin:0 0 22px}.login-card .field{margin-bottom:14px}.logout{display:block;color:#a7b0c0;text-decoration:none;font-size:13px;margin:22px 10px 0}
textarea{width:100%;min-height:220px;border:1px solid var(--line);border-radius:8px;padding:12px;font-family:Consolas,monospace;resize:vertical}
table{width:100%;border-collapse:collapse;background:white;border:1px solid var(--line);border-radius:8px;overflow:hidden}th,td{text-align:left;border-bottom:1px solid var(--line);padding:8px 9px;font-size:13px;vertical-align:top}th{background:#f8fafc}td{word-break:break-word}.status{padding:10px 12px;border-left:4px solid var(--accent);background:#eef7f5;border-radius:0 8px 8px 0;margin:12px 0;display:none}
.progress{display:none;margin-top:12px}.progress .track{height:12px;background:#e5e7eb;border-radius:999px;overflow:hidden}.progress .bar{width:0;height:100%;background:var(--accent);transition:width .2s ease}.progress.active .bar{background:linear-gradient(90deg,var(--accent),#22c55e,var(--accent));background-size:180% 100%;animation:slide 1.2s linear infinite}.progress .label{font-size:13px;color:var(--muted);margin-top:7px}@keyframes slide{from{background-position:0 0}to{background-position:180% 0}}
@media(max-width:900px){.app{grid-template-columns:1fr}.grid,.cards{grid-template-columns:1fr}.side{position:relative}}
</style>
"""


def shell(title: str, active: str, main: str) -> bytes:
    icons = {
        "upload": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M12 16V4"></path><path d="m7 9 5-5 5 5"></path><path d="M20 16v4H4v-4"></path></svg>',
        "organizer": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M3 6h7l2 2h9v10a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2Z"></path><path d="M8 13h8"></path><path d="M8 16h5"></path></svg>',
        "grid": '<svg viewBox="0 0 24 24" aria-hidden="true"><rect x="3" y="3" width="7" height="7"></rect><rect x="14" y="3" width="7" height="7"></rect><rect x="3" y="14" width="7" height="7"></rect><rect x="14" y="14" width="7" height="7"></rect></svg>',
        "urls": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M10 13a5 5 0 0 0 7.1 0l2-2a5 5 0 0 0-7.1-7.1l-1.1 1.1"></path><path d="M14 11a5 5 0 0 0-7.1 0l-2 2A5 5 0 0 0 12 20.1l1.1-1.1"></path></svg>',
        "excel": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8Z"></path><path d="M14 2v6h6"></path><path d="M8 13h8"></path><path d="M8 17h5"></path></svg>',
        "riegel": '<svg viewBox="0 0 24 24" aria-hidden="true"><circle cx="11" cy="11" r="7"></circle><path d="m20 20-3.5-3.5"></path><path d="M8 11h6"></path><path d="M11 8v6"></path></svg>',
        "gewicht": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M4 20h16"></path><path d="M6 20 9 8h6l3 12"></path><circle cx="12" cy="5" r="2"></circle></svg>',
    }
    nav = f"""
      <div class="brand-wrap">
        <div class="logo-row">
          <img class="brand-logo primex" src="/assets/logoprimexeu.png" alt="Primex EU logo">
        </div>
      </div>
      <nav class="nav">
        <a class="{'active' if active == 'upload' else ''}" href="/upload">{icons['upload']}<span>Upload Images</span></a>
        <a class="{'active' if active == 'organizer' else ''}" href="/photo-organizer">{icons['organizer']}<span>Photo Organizer</span></a>
        <a class="{'active' if active == 'grid' else ''}" href="/image-generator">{icons['grid']}<span>Image Grid Generator</span></a>
        <a class="{'active' if active == 'urls' else ''}" href="/url-generator">{icons['urls']}<span>URL Generator</span></a>
        <a class="{'active' if active == 'excel' else ''}" href="/excel-url-filler">{icons['excel']}<span>Excel URL Filler</span></a>
        <a class="{'active' if active == 'riegel' else ''}" href="/riegel-logo-scanner">{icons['riegel']}<span>Riegel Logo Scanner</span></a>
        <a class="{'active' if active == 'gewicht' else ''}" href="/gewicht">{icons['gewicht']}<span>GEWICHT</span></a>
      </nav>
      <a class="logout" href="/logout">Sign out</a>
    """
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title><link rel="icon" type="image/png" href="/assets/logoprimexeu.png">{STYLE}</head><body><div class="app"><aside class="side">{nav}</aside><main>{main}</main></div>{SCRIPT}</body></html>""".encode("utf-8")


LOGIN_PAGE = """
<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>VS Tools Login</title><link rel="icon" type="image/png" href="/assets/logoprimexeu.png">__STYLE__</head>
<body><div class="login-page"><form class="login-card" method="post" action="/login">
  <img class="login-logo" src="/assets/logoprimexeu.png" alt="Primex EU">
  <div class="field"><label>Email</label><input name="email" type="email" autocomplete="username" required></div>
  <div class="field"><label>Password</label><input name="password" type="password" autocomplete="current-password" required></div>
  <button type="submit">Sign in</button>
  <div class="status" id="status" style="display:__ERROR_DISPLAY__">__ERROR__</div>
</form></div></body></html>
""".replace("__STYLE__", STYLE)


URL_PAGE = """
<h1>Generate Image URLs</h1>
<p class="hint">Upload folders to Hetzner as usual. This page detects them automatically; choose a folder and generate URLs for all images inside it.</p>
<section class="panel">
  <div class="grid">
    <div class="field"><label>Search top folders</label><input id="topSearch" placeholder="Type to filter recent top folders"></div>
    <div class="field"><label>Search product folders</label><input id="productSearch" placeholder="Type to filter products"></div>
    <div class="field"><label>URL base</label><input value="__BASE_URL__" readonly></div>
    <div class="field"><label>Top folder</label><select id="topFolder"></select></div>
    <div class="field"><label>Product folder</label><select id="productFolder"><option value="">All product folders</option></select></div>
    <div class="field"><label>Recent folders</label><select id="recentFolder"><option value="">Choose recent top folder</option></select></div>
    <div class="field full"><div class="checks"><label class="check"><input type="checkbox" id="recursive" checked> Include subfolders recursively</label><label class="check"><input type="checkbox" id="cacheBust"> Add cache version to URLs</label><label class="check"><input type="checkbox" id="showThumbs"> Show thumbnails in table</label></div></div>
  </div>
  <div class="actions">
    <button id="generateBtn">Get URLs</button><button class="secondary" id="checkUrlBtn">Check URLs</button><button class="secondary" id="refreshBtn">Refresh Folders</button><button class="mutedbtn" id="copyBtn">Copy URLs</button><button class="mutedbtn" id="downloadBtn">Download TXT</button>
  </div>
  <div class="status" id="status"></div>
</section>
<section class="panel"><div class="cards"><div class="card"><div class="n" id="topCount">0</div><div class="t">top folders detected</div></div><div class="card"><div class="n" id="productCount">0</div><div class="t">product folders in selected folder</div></div><div class="card"><div class="n" id="urlCount">0</div><div class="t">URLs generated</div></div><div class="card"><div class="n" id="onlineCount">0</div><div class="t">online URLs</div></div><div class="card"><div class="n" id="errorCount">0</div><div class="t">URL errors</div></div><div class="card"><div class="n" id="latestFolder">-</div><div class="t">latest top folder</div></div></div></section>
<section class="panel"><label>Generated URLs</label><textarea id="urlBox" readonly placeholder="URLs will appear here..."></textarea></section>
<section class="panel"><h2>Files</h2><div class="toolbar"><div class="field"><label>Problem filter</label><select id="resultFilter"><option value="all">Show all</option><option value="errors">Only errors</option><option value="missing">Only missing source</option><option value="skipped">Only skipped</option></select></div><button class="mutedbtn" id="exportErrorsBtn">Export Errors</button></div><div id="tableWrap" class="hint">No URLs generated yet.</div></section>
""".replace("__BASE_URL__", html.escape(BASE_URL))


GRID_PAGE = """
<h1>Image Grid Generator</h1>
<p class="hint">Select a Hetzner folder and generate 1, 3, 6, 12, or 12-alt grid images directly in each product folder.</p>
<section class="panel">
  <div class="grid">
    <div class="field"><label>Search top folders</label><input id="topSearch" placeholder="Type to filter recent top folders"></div>
    <div class="field"><label>Search product folders</label><input id="productSearch" placeholder="Type to filter products"></div>
    <div class="field"><label>Recent folders</label><select id="recentFolder"><option value="">Choose recent top folder</option></select></div>
    <div class="field"><label>Top folder</label><select id="topFolder"></select></div>
    <div class="field"><label>Product folder</label><select id="productFolder"><option value="">All product folders</option></select></div>
    <div class="field"><label>Format</label><select id="format"><option value="jpg">JPG</option><option value="png">PNG</option></select></div>
    <div class="field full"><label>Quantities</label><div class="checks"><label class="check"><input type="checkbox" class="qty" value="1" checked> 1</label><label class="check"><input type="checkbox" class="qty" value="3" checked> 3</label><label class="check"><input type="checkbox" class="qty" value="6" checked> 6</label><label class="check"><input type="checkbox" class="qty" value="12" checked> 12</label><label class="check"><input type="checkbox" class="qty" value="12_alt"> 12 alt</label><label class="check"><input type="checkbox" id="deleteSource"> Delete source after successful generation</label><label class="check"><input type="checkbox" id="gridCacheBust"> Add cache version to URLs</label><label class="check"><input type="checkbox" id="showThumbs"> Show thumbnails in table</label></div></div>
  </div>
  <div class="actions"><button id="gridBtn">Generate Images</button><button class="secondary" id="refreshBtn">Refresh Folders</button></div>
  <div class="status" id="status"></div>
</section>
<section class="panel"><h2>Results</h2><div class="toolbar"><div class="field"><label>Problem filter</label><select id="resultFilter"><option value="all">Show all</option><option value="errors">Only errors</option><option value="missing">Only missing source</option><option value="skipped">Only skipped</option></select></div><button class="mutedbtn" id="exportErrorsBtn">Export Errors</button></div><div id="tableWrap" class="hint">No images generated yet.</div></section>
"""


ORGANIZER_PAGE = """
<h1>Photo Organizer</h1>
<p class="hint">Organize loose images from a top folder into product-code subfolders. The code is taken from the filename before the first underscore.</p>
<section class="panel">
  <div class="grid">
    <div class="field"><label>Search source folders</label><input id="topSearch" placeholder="Type to filter recent folders"></div>
    <div class="field"><label>Recent folders</label><select id="recentFolder"><option value="">Choose recent top folder</option></select></div>
    <div class="field"><label>Source top folder</label><select id="topFolder"></select></div>
    <div class="field"><label>Destination top folder</label><input id="destinationFolder" placeholder="Example: WW2"></div>
    <div class="field"><label>Mode</label><select id="organizeMode"><option value="copy">Copy files</option><option value="move">Move files</option></select></div>
  </div>
  <div class="actions"><button id="organizeBtn">Organize Photos</button><button class="secondary" id="refreshBtn">Refresh Folders</button></div>
  <div class="status" id="status"></div>
</section>
<section class="panel"><h2>Results</h2><div id="tableWrap" class="hint">No files organized yet.</div></section>
"""


UPLOAD_PAGE = """
<h1>Upload Images</h1>
<p class="hint">Upload directly into the Hetzner photos folder. The fixed base is <strong>/var/www/primexeu.com.shared/photos/VS</strong>; enter only the path after <strong>VS/</strong>.</p>
<section class="panel">
  <div class="grid">
    <div class="field full"><label>Destination path after VS/</label><input id="uploadPath" list="uploadPathSuggestions" autocomplete="off" placeholder="Example: FG or FG/PFH511 or WW2/001374"><datalist id="uploadPathSuggestions"></datalist></div>
    <div class="field full"><label>Files</label><input id="uploadFiles" type="file" multiple accept=".jpg,.jpeg,.png,.webp,.gif,.pdf"></div>
    <div class="field full"><div class="checks"><label class="check"><input type="checkbox" id="overwriteFiles"> Overwrite existing files</label><label class="check"><input type="checkbox" id="checkYears"> Check for year on bottle</label><label class="check"><input type="checkbox" id="automateUpload" checked> After upload, automatically organize, generate grid images, and return URLs</label></div></div>
    <div class="field"><label>Grid format</label><select id="autoFormat"><option value="jpg">JPG</option><option value="png">PNG</option></select></div>
    <div class="field full"><label>Grid quantities for automation</label><div class="checks"><label class="check"><input type="checkbox" class="autoQty" value="1" checked> 1</label><label class="check"><input type="checkbox" class="autoQty" value="3" checked> 3</label><label class="check"><input type="checkbox" class="autoQty" value="6" checked> 6</label><label class="check"><input type="checkbox" class="autoQty" value="12" checked> 12</label><label class="check"><input type="checkbox" class="autoQty" value="12_alt"> 12 alt</label><label class="check"><input type="checkbox" id="autoDeleteSource" checked> Delete source after grid generation</label><label class="check"><input type="checkbox" id="autoCacheBust"> Add cache version to URLs</label></div></div>
  </div>
  <div class="actions"><button id="uploadBtn">Upload Files</button><button class="secondary" id="refreshBtn">Refresh Folders</button></div>
  <div class="progress" id="uploadProgress"><div class="track"><div class="bar" id="uploadProgressBar"></div></div><div class="label" id="uploadProgressLabel">Waiting...</div></div>
  <div class="status" id="status"></div>
</section>
<section class="panel"><h2>Upload Results</h2><div class="actions"><button class="mutedbtn" id="copyUploadLinksBtn">Copy All Links</button><button class="mutedbtn" id="downloadUploadLinksBtn">Export TXT</button></div><div id="tableWrap" class="hint">No files uploaded yet.</div></section>
"""


EXCEL_PAGE = """
<h1>Excel URL Filler</h1>
<p class="hint">Upload the Amazon .xlsm after image URLs are generated. This fills column <strong>URL des Hauptbildes</strong> and the first <strong>Andere Bild-URL</strong> from the Hetzner folder.</p>
<section class="panel">
  <div class="grid">
    <div class="field"><label>Hetzner top folder after VS/</label><input id="excelTopFolder" value="KR" list="excelTopFolderSuggestions" autocomplete="off" placeholder="Example: KR"><datalist id="excelTopFolderSuggestions"></datalist></div>
    <div class="field"><label>SKU prefix to remove</label><input id="excelPrefix" value="KR" placeholder="Example: KR"></div>
    <div class="field"><label>SKU suffix to remove</label><input id="excelSuffix" value="bio" placeholder="Example: bio"></div>
    <div class="field full"><label>Amazon Excel file</label><input id="excelFile" type="file" accept=".xlsm,.xlsx"></div>
    <div class="field full"><div class="checks"><label class="check"><input type="checkbox" id="excelCacheBust"> Add cache version to URLs</label></div></div>
  </div>
  <div class="actions"><button id="excelBtn">Fill Excel URLs</button></div>
  <div class="progress" id="excelProgress"><div class="track"><div class="bar" id="excelProgressBar"></div></div><div class="label" id="excelProgressLabel">Waiting...</div></div>
  <div class="status" id="status"></div>
</section>
<section class="panel"><h2>Rules Used</h2><div class="hint">Example: <strong>03-KRN24bio</strong> becomes folder <strong>KR/N24</strong>, main image <strong>3.jpg</strong>, and other image <strong>1.jpg</strong>.</div></section>
"""


GEWICHT_PAGE = """
<h1>GEWICHT</h1>
<p class="hint">Ngarko një Excel me dy kolonat <strong>Artikel Nr</strong> dhe <strong>Gewicht</strong>. Shkarko Excel-in e ri me llogaritjen për çdo rresht.</p>
<section class="panel">
  <div class="field"><label for="gewichtFile">Excel (.xlsx ose .xlsm)</label><input id="gewichtFile" type="file" accept=".xlsx,.xlsm"></div>
  <div class="field"><label for="gewichtPrefix">Prefiksi i kodit të artikullit</label><input id="gewichtPrefix" value="AM" placeholder="P.sh. AM" maxlength="20" required></div>
  <div class="field"><label for="gewichtUnit">Njësia e peshës në Excel</label><select id="gewichtUnit"><option value="g">Gramë (g) — kthe në mg</option><option value="mg">Miligramë (mg) — pa konvertim</option></select></div>
  <div class="field full"><label for="gewichtCodeSuffix">Prapashtesë e kodit të artikullit (opsionale)</label><input id="gewichtCodeSuffix" placeholder="P.sh. BIO → 01-AM5657BIO; lëre bosh për kodin pa prapashtesë" maxlength="20"></div>
  <div class="actions"><button id="gewichtBtn">Llogarit dhe shkarko Excel</button></div>
  <div class="status" id="status" role="status"></div>
</section>
<section class="panel">
  <h2>Rregullat e llogaritjes</h2>
  <p class="hint">Zgjidh g ose mg para ngarkimit. Kur zgjedh g, pesha origjinale ruhet dhe shtohet kolona e konvertuar në mg (1 g = 1000 mg). Kur zgjedh mg, nuk bëhet konvertim. Vlerat me njësi të shkruar brenda qelizës lexohen sipas asaj njësie.</p>
  <p class="hint">Numri i shisheve përsëritet 1, 3, 6, 12. Kartoni përkatës është 180, 540, 1120, 2120 mg. GEWICHT = pesha në mg × numri i shisheve + kartoni. Kodi krijohet si 01-AM5657, 03-AM5657, 06-AM5657, 12-AM5657, sipas prefiksit të zgjedhur.</p>
</section>
"""


RIEGEL_PAGE = """
<h1>Riegel Logo Scanner</h1>
<p class="hint">Scan image URLs or a local/network folder for visible Riegel logo text. This report only identifies matches; it does not remove or edit images.</p>
<section class="panel">
  <div class="grid">
    <div class="field full"><label>Image URLs</label><textarea id="riegelUrls" placeholder="Paste one image URL per line..."></textarea></div>
    <div class="field full"><label>Folder path</label><input id="riegelFolder" placeholder="Example: C:\\Users\\Admin\\Downloads or W:\\05_CLIENTS\\02_VS\\10_DESIGN\\01_WINE-SPIRITS\\25_Riegel_V_PROJECT\\ALL IMAGES"></div>
    <div class="field full"><div class="checks"><label class="check"><input type="checkbox" id="riegelRecursive" checked> Include subfolders</label></div></div>
  </div>
  <div class="actions"><button id="riegelScanBtn">Scan for Riegel Logo</button><button class="mutedbtn" id="copyRiegelHitsBtn">Copy Found Paths / URLs</button><button class="mutedbtn" id="exportRiegelHitsBtn">Export Found CSV</button></div>
  <div class="progress" id="riegelProgress"><div class="track"><div class="bar" id="riegelProgressBar"></div></div><div class="label" id="riegelProgressLabel">Waiting...</div></div>
  <div class="status" id="status"></div>
</section>
<section class="panel">
  <div class="cards"><div class="card"><div class="n" id="riegelScannedCount">0</div><div class="t">images scanned</div></div><div class="card"><div class="n" id="riegelFoundCount">0</div><div class="t">Riegel found</div></div><div class="card"><div class="n" id="riegelErrorCount">0</div><div class="t">errors/skips</div></div></div>
</section>
<section class="panel"><h2>Scan Results</h2><div id="tableWrap" class="hint">No scan run yet.</div></section>
"""


SCRIPT = """
<script>
let currentRows = [];
let topFolderCache = [];
let currentProducts = [];
const productFolderCache = {};
function status(msg){const el=document.getElementById("status"); if(el){el.textContent=msg; el.style.display="block";}}
async function api(path){const res=await fetch(path,{cache:"no-store"}); const data=await res.json(); if(!res.ok) throw new Error(data.error||"Request failed"); return data;}
function matchesSearch(text, query){return !query || String(text||"").toLowerCase().includes(query.toLowerCase());}
function renderTopFolders(keep){const select=document.getElementById("topFolder"); if(!select)return; const old=keep||select.value; const query=document.getElementById("topSearch")?.value||""; const filtered=topFolderCache.filter(folder=>matchesSearch(folder.name,query)); select.innerHTML=""; filtered.forEach(folder=>{const opt=document.createElement("option"); opt.value=folder.name; opt.textContent=`${folder.name} (${folder.subfolderCount} folders, ${folder.modifiedText})`; select.appendChild(opt);}); if(old&&[...select.options].some(o=>o.value===old))select.value=old; const recent=document.getElementById("recentFolder"); if(recent){const recentOld=recent.value; recent.innerHTML='<option value="">Choose recent top folder</option>'+topFolderCache.slice(0,30).map(folder=>`<option value="${escAttr(folder.name)}">${escAttr(folder.name)} (${escAttr(folder.modifiedText)})</option>`).join(""); if(recentOld&&[...recent.options].some(o=>o.value===recentOld))recent.value=recentOld;}}
async function loadTopFolders(keep){const data=await api("/api/top-folders"); topFolderCache=data.folders||[]; updateUploadPathSuggestions(); updateExcelTopFolderSuggestions(); renderTopFolders(keep); const tc=document.getElementById("topCount"); if(tc)tc.textContent=data.folders.length; const lf=document.getElementById("latestFolder"); if(lf)lf.textContent=data.folders[0]?.name||"-"; await loadProducts(); status(`Detected ${data.folders.length} top folders. Last refreshed ${new Date().toLocaleTimeString()}.`);}
function renderProducts(keep){const select=document.getElementById("productFolder"); if(!select)return; const old=keep||select.value; const query=document.getElementById("productSearch")?.value||""; select.innerHTML='<option value="">All product folders</option>'; currentProducts.filter(product=>matchesSearch(product.name,query)).forEach(product=>{const opt=document.createElement("option"); opt.value=product.name; opt.textContent=`${product.name} (${product.imageCount} files, ${product.modifiedText})`; select.appendChild(opt);}); if(old&&[...select.options].some(o=>o.value===old))select.value=old;}
async function loadProducts(){const top=document.getElementById("topFolder")?.value; const select=document.getElementById("productFolder"); if(!select)return; currentProducts=[]; if(top){const data=await api("/api/products?folder="+encodeURIComponent(top)); currentProducts=data.products||[];} renderProducts(); const pc=document.getElementById("productCount"); if(pc)pc.textContent=currentProducts.length;}
function escAttr(value){return String(value).replace(/[&<>"']/g,ch=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[ch]));}
async function updateUploadPathSuggestions(){const input=document.getElementById("uploadPath"); const list=document.getElementById("uploadPathSuggestions"); if(!input||!list)return; const value=input.value.trim().replace(/\\\\/g,"/"); const parts=value.split("/").filter(Boolean); let suggestions=[]; if(value.includes("/")&&parts[0]){const top=parts[0]; if(!productFolderCache[top]){const data=await api("/api/products?folder="+encodeURIComponent(top)); productFolderCache[top]=data.products||[];} suggestions=productFolderCache[top].map(product=>`${top}/${product.name}`);}else{suggestions=topFolderCache.map(folder=>folder.name); suggestions.push(...topFolderCache.map(folder=>`${folder.name}/`));} list.innerHTML=[...new Set(suggestions)].slice(0,200).map(item=>`<option value="${escAttr(item)}"></option>`).join("");}
function updateExcelTopFolderSuggestions(){const list=document.getElementById("excelTopFolderSuggestions"); if(!list)return; list.innerHTML=topFolderCache.slice(0,250).map(folder=>`<option value="${escAttr(folder.name)}"></option>`).join("");}
function resetUrlCheckCounts(){const online=document.getElementById("onlineCount"); const error=document.getElementById("errorCount"); if(online)online.textContent="0"; if(error)error.textContent="0";}
function showThumbs(){return document.getElementById("showThumbs")?.checked===true;}
function isImageUrl(url){return /\\.(jpe?g|png|webp|gif)(\\?|#|$)/i.test(String(url||""));}
function thumbCell(row){if(!showThumbs())return ""; const url=row.url||""; if(!isImageUrl(url))return '<td class="thumb-cell"><span class="hint">No preview</span></td>'; return `<td class="thumb-cell"><a href="${url}" target="_blank"><img class="thumb" src="${url}" loading="lazy" alt=""></a></td>`;}
function problemType(row){const status=String(row.status||"").toLowerCase(); const message=String(row.message||"").toLowerCase(); if(row.ok===false||status.includes("error"))return "errors"; if(status.includes("missing")||message.includes("missing")||message.includes("no source"))return "missing"; if(status.includes("skipped"))return "skipped"; return "ok";}
function filteredRows(){const filter=document.getElementById("resultFilter")?.value||"all"; if(filter==="all")return currentRows; return currentRows.filter(row=>problemType(row)===filter);}
function renderCurrentTable(){if(document.getElementById("generateBtn"))renderUrlTable(); else if(document.getElementById("gridBtn"))renderGridTable(); else if(document.getElementById("uploadBtn"))renderUploadTable(); else if(document.getElementById("organizeBtn"))renderOrganizeTable(); else if(document.getElementById("riegelScanBtn"))renderRiegelTable();}
function exportProblemRows(){const rows=currentRows.filter(row=>problemType(row)!=="ok"); if(!rows.length){status("No problem rows to export.");return;} const lines=rows.map(row=>[row.step||"",row.product||"",row.file||row.source||row.created||"",row.status||"",row.message||"",row.path||"",row.url||""].map(value=>`"${String(value).replace(/"/g,'""')}"`).join(",")); const blob=new Blob([["step","product","file","status","message","path","url"].join(","),...lines].join("\\n"),{type:"text/csv"}); const a=document.createElement("a"); a.href=URL.createObjectURL(blob); a.download="primex-problem-files.csv"; a.click(); URL.revokeObjectURL(a.href); status(`Exported ${rows.length} problem row(s).`);}
async function generateUrls(){const top=document.getElementById("topFolder").value; const product=document.getElementById("productFolder").value; const recursive=document.getElementById("recursive").checked?"1":"0"; const cache=document.getElementById("cacheBust")?.checked?"1":"0"; const data=await api(`/api/generate?folder=${encodeURIComponent(top)}&product=${encodeURIComponent(product)}&recursive=${recursive}&cache=${cache}`); currentRows=data.rows; document.getElementById("urlBox").value=currentRows.map(row=>row.url).join("\\n"); document.getElementById("urlCount").textContent=currentRows.length; resetUrlCheckCounts(); renderUrlTable(); status(`Generated ${currentRows.length} URLs from ${top}${product?"/"+product:""}.`);}
async function checkUrls(){if(!currentRows.length)await generateUrls(); if(!currentRows.length){status("No URLs to check.");return;} const btn=document.getElementById("checkUrlBtn"); if(btn)btn.disabled=true; currentRows.forEach(row=>{row.ok=null; row.status="Waiting";}); renderUrlTable(); let online=0; let errors=0; try{for(let i=0;i<currentRows.length;i++){const row=currentRows[i]; row.status="Checking..."; renderUrlRow(i); status(`Checking URLs: ${i+1} / ${currentRows.length} (${online} online, ${errors} errors)`); try{const data=await api(`/api/check-url?url=${encodeURIComponent(row.url)}`); row.ok=data.ok; row.status=data.status; if(data.ok)online++; else errors++;}catch(err){row.ok=false; row.status=err.message||"ERROR"; errors++;} document.getElementById("onlineCount").textContent=online; document.getElementById("errorCount").textContent=errors; renderUrlRow(i); await new Promise(resolve=>setTimeout(resolve,0));} status(`Checked ${currentRows.length} URLs: ${online} online, ${errors} error(s).`);}finally{if(btn)btn.disabled=false;}}
function urlStatusCell(row){if(row.status==="Checking...")return '<strong style="color:#1d4ed8">CHECKING</strong>'; if(row.status==="Waiting")return '<span class="hint">Waiting</span>'; if(!row.status)return '<span class="hint">Not checked</span>'; const ok=row.ok===true; const color=ok?"#0f766e":"#b42318"; const label=ok?"ONLINE":"ERROR"; return `<strong style="color:${color}">${label}</strong><div class="hint">${row.status}</div>`;}
function urlRowHtml(row,index){return `<tr data-url-row="${index}">${thumbCell(row)}<td>${row.product||""}</td><td>${row.file}</td><td>${row.path}</td><td>${urlStatusCell(row)}</td><td><a href="${row.url}" target="_blank">${row.url}</a></td></tr>`;}
function renderUrlRow(index){const tr=document.querySelector(`[data-url-row="${index}"]`); if(tr)tr.outerHTML=urlRowHtml(currentRows[index],index);}
function renderUrlTable(){if(!currentRows.length){document.getElementById("tableWrap").textContent="No URLs generated.";return;} const visible=filteredRows(); const rows=visible.slice(0,500).map(row=>urlRowHtml(row,currentRows.indexOf(row))).join(""); const note=visible.length>500?`<p class="hint">Showing first 500 of ${visible.length} rows.</p>`:""; const preview=showThumbs()?"<th>Preview</th>":""; document.getElementById("tableWrap").innerHTML=`${note}<table><thead><tr>${preview}<th>Product</th><th>File</th><th>Path</th><th>Status</th><th>URL</th></tr></thead><tbody>${rows}</tbody></table>`;}
async function generateGrid(){const top=document.getElementById("topFolder").value; const product=document.getElementById("productFolder").value; const fmt=document.getElementById("format").value; const del=document.getElementById("deleteSource").checked?"1":"0"; const cache=document.getElementById("gridCacheBust")?.checked?"1":"0"; const qty=[...document.querySelectorAll(".qty:checked")].map(q=>q.value).join(","); status("Generating images. Keep this page open..."); const data=await api(`/api/generate-images?folder=${encodeURIComponent(top)}&product=${encodeURIComponent(product)}&format=${fmt}&quantities=${encodeURIComponent(qty)}&delete=${del}&cache=${cache}`); currentRows=data.rows; renderGridTable(); status(`Finished. Created ${data.created} images. Errors/skips: ${data.issues}.`);}
function renderGridTable(){if(!currentRows.length){document.getElementById("tableWrap").textContent="No results.";return;} const visible=filteredRows(); const rows=visible.slice(0,800).map(row=>`<tr>${thumbCell(row)}<td>${row.product||""}</td><td>${row.source||""}</td><td>${row.created||""}</td><td>${row.status||""}</td><td>${row.url?`<a href="${row.url}" target="_blank">${row.url}</a>`:(row.message||"")}</td></tr>`).join(""); const preview=showThumbs()?"<th>Preview</th>":""; document.getElementById("tableWrap").innerHTML=`<table><thead><tr>${preview}<th>Product</th><th>Source</th><th>Created</th><th>Status</th><th>URL / Message</th></tr></thead><tbody>${rows}</tbody></table>`;}
async function organizePhotos(){const source=document.getElementById("topFolder").value; const destination=document.getElementById("destinationFolder").value || source; const move=document.getElementById("organizeMode").value==="move"?"1":"0"; status("Organizing photos..."); const data=await api(`/api/organize-photos?source=${encodeURIComponent(source)}&destination=${encodeURIComponent(destination)}&move=${move}`); currentRows=data.rows; renderOrganizeTable(); status(`Finished. ${data.changed} files copied/moved. Issues/skips: ${data.issues}.`);}
function renderOrganizeTable(){if(!currentRows.length){document.getElementById("tableWrap").textContent="No results.";return;} const rows=currentRows.slice(0,800).map(row=>`<tr><td>${row.file||""}</td><td>${row.product||""}</td><td>${row.status||""}</td><td>${row.url?`<a href="${row.url}" target="_blank">${row.url}</a>`:(row.message||row.target||"")}</td></tr>`).join(""); document.getElementById("tableWrap").innerHTML=`<table><thead><tr><th>File</th><th>Product</th><th>Status</th><th>URL / Message</th></tr></thead><tbody>${rows}</tbody></table>`;}
function setUploadProgress(percent,label,active=false){const wrap=document.getElementById("uploadProgress"); const bar=document.getElementById("uploadProgressBar"); const text=document.getElementById("uploadProgressLabel"); if(!wrap||!bar||!text)return; wrap.style.display="block"; wrap.classList.toggle("active",active); bar.style.width=`${Math.max(0,Math.min(100,percent))}%`; text.textContent=label;}
async function pollUploadJob(jobId,priorRows=[]){let lastRows=0; while(true){const data=await api(`/api/job?id=${encodeURIComponent(jobId)}`); currentRows=[...priorRows,...(data.rows||[])]; if(currentRows.length!==lastRows || data.done){renderUploadTable(); lastRows=currentRows.length;} const label=data.message||data.phase||"Working..."; setUploadProgress(Math.max(75,Number(data.percent||75)),label,!data.done); status(`${label} Uploaded ${data.uploaded||0}. Created ${data.created||0} grid image(s). URLs ${data.urls||0}. Year images ${data.years||0}. Issues/skips ${data.issues||0}.`); if(data.done){setUploadProgress(100,data.ok===false?"Finished with errors.":"Finished.",false); renderUploadTable(); return data;} await new Promise(resolve=>setTimeout(resolve,1200));}}
async function uploadFiles(){
  const path=document.getElementById("uploadPath").value;
  const files=[...document.getElementById("uploadFiles").files];
  const automate=document.getElementById("automateUpload").checked;
  const checkYears=document.getElementById("checkYears").checked;
  const btn=document.getElementById("uploadBtn");
  if(!path.trim()){status("Enter destination path after VS/.");return;}
  if(!files.length){status("Choose files to upload.");return;}
  const maxBatchBytes=75*1024*1024;
  const maxFileBytes=95*1024*1024;
  const tooLarge=files.find(file=>file.size>maxFileBytes);
  if(tooLarge){status(`${tooLarge.name} is larger than 95 MB. Upload it through a direct/DNS-only route or reduce its size.`);return;}
  const batches=[];
  let batch=[];
  let batchBytes=0;
  for(const file of files){if(batch.length&&batchBytes+file.size>maxBatchBytes){batches.push(batch);batch=[];batchBytes=0;} batch.push(file);batchBytes+=file.size;}
  if(batch.length)batches.push(batch);
  const totalBytes=Math.max(1,files.reduce((sum,file)=>sum+file.size,0));
  let completedBytes=0;
  let allRows=[];
  let uploaded=0;
  let years=0;
  let issues=0;
  const commonFields=form=>{form.append("path",path);form.append("overwrite",document.getElementById("overwriteFiles").checked?"1":"0");form.append("format",document.getElementById("autoFormat").value);form.append("delete",document.getElementById("autoDeleteSource").checked?"1":"0");form.append("cache_bust",document.getElementById("autoCacheBust")?.checked?"1":"0");form.append("quantities",[...document.querySelectorAll(".autoQty:checked")].map(q=>q.value).join(","));};
  const sendForm=(form,onProgress)=>new Promise((resolve,reject)=>{const xhr=new XMLHttpRequest();xhr.open("POST","/api/upload");xhr.timeout=30*60*1000;xhr.upload.onprogress=onProgress;xhr.onload=()=>{let payload={};try{payload=JSON.parse(xhr.responseText||"{}");}catch(err){reject(new Error(xhr.status===413?"Upload batch is too large.":`Upload returned invalid response (HTTP ${xhr.status}).`));return;}if(xhr.status>=200&&xhr.status<300)resolve(payload);else reject(new Error(payload.error||`Upload failed (HTTP ${xhr.status})`));};xhr.onerror=()=>reject(new Error("Upload connection failed."));xhr.ontimeout=()=>reject(new Error("Upload timed out after 30 minutes."));xhr.send(form);});
  if(btn)btn.disabled=true;
  try{
    status(`Uploading ${files.length} file(s) in ${batches.length} safe batch(es)...`);
    setUploadProgress(1,"Preparing upload...");
    for(let index=0;index<batches.length;index++){
      const current=batches[index];
      const currentBytes=current.reduce((sum,file)=>sum+file.size,0);
      const form=new FormData();commonFields(form);form.append("check_years",checkYears?"1":"0");form.append("automate","0");current.forEach(file=>form.append("files",file));
      const data=await sendForm(form,event=>{if(!event.lengthComputable)return;const sent=completedBytes+(event.loaded/event.total)*currentBytes;const pct=Math.max(1,Math.round((sent/totalBytes)*(automate?72:96)));setUploadProgress(pct,`Uploading batch ${index+1}/${batches.length}: ${Math.round((sent/totalBytes)*100)}% overall`);});
      completedBytes+=currentBytes;
      allRows.push(...(data.rows||[]));uploaded+=Number(data.uploaded||0);years+=Number(data.years||0);issues+=Number(data.issues||0);currentRows=allRows;renderUploadTable();
    }
    if(automate&&uploaded){
      setUploadProgress(74,"Uploads finished. Starting automation...",true);
      const form=new FormData();commonFields(form);form.append("check_years","0");form.append("automate","1");form.append("automate_only","1");form.append("uploaded_count",String(uploaded));
      const data=await sendForm(form,()=>{});
      if(data.job_id){await pollUploadJob(data.job_id,allRows);return;}
    }
    setUploadProgress(100,"Finished.",false);
    status(`Uploaded ${uploaded} file(s). Year detected in ${years} image(s). Issues/skips: ${issues}.`);
  }finally{if(btn)btn.disabled=false;}
}
function setExcelProgress(percent,label,active=false){const wrap=document.getElementById("excelProgress"); const bar=document.getElementById("excelProgressBar"); const text=document.getElementById("excelProgressLabel"); if(!wrap||!bar||!text)return; wrap.style.display="block"; wrap.classList.toggle("active",active); bar.style.width=`${Math.max(0,Math.min(100,percent))}%`; text.textContent=label;}
async function fillExcelUrls(){const file=document.getElementById("excelFile").files[0]; if(!file){status("Choose an Excel file.");return;} const form=new FormData(); form.append("file",file); form.append("top",document.getElementById("excelTopFolder").value); form.append("prefix",document.getElementById("excelPrefix").value); form.append("suffix",document.getElementById("excelSuffix").value); form.append("cache_bust",document.getElementById("excelCacheBust")?.checked?"1":"0"); status("Uploading Excel file..."); setExcelProgress(5,"Uploading Excel file..."); const result=await new Promise((resolve,reject)=>{const xhr=new XMLHttpRequest(); xhr.open("POST","/api/fill-excel"); xhr.responseType="blob"; xhr.upload.onprogress=(event)=>{if(event.lengthComputable)setExcelProgress(Math.round((event.loaded/event.total)*50),`Uploading Excel: ${Math.round((event.loaded/event.total)*100)}%`);}; xhr.upload.onload=()=>setExcelProgress(65,"Server is filling columns and checking URLs...",true); xhr.onload=()=>{if(xhr.status>=200&&xhr.status<300){let stats={}; try{stats=JSON.parse(xhr.getResponseHeader("X-Excel-Url-Stats")||"{}");}catch(err){} resolve({blob:xhr.response,stats});return;} const reader=new FileReader(); reader.onload=()=>{try{reject(new Error(JSON.parse(reader.result).error||"Excel fill failed"));}catch(err){reject(new Error("Excel fill failed"));}}; reader.readAsText(xhr.response);}; xhr.onerror=()=>reject(new Error("Excel upload connection failed.")); xhr.send(form);}); const a=document.createElement("a"); a.href=URL.createObjectURL(result.blob); const baseName=file.name.replace(/^(?:filled-)+/i,""); a.download=`filled-${baseName}`; a.click(); URL.revokeObjectURL(a.href); const stats=result.stats||{}; const checked=Number(stats.checked||0); const invalid=Number(stats.invalid||0); const valid=Number(stats.valid||0); const missingProducts=Array.isArray(stats.missing_products)?stats.missing_products:[]; const missingText=missingProducts.length?` Missing products: ${missingProducts.join(", ")}.`:""; setExcelProgress(100,"Finished. Download started.",false); status(checked?`Excel URLs filled. Checked ${checked} URL(s): ${valid} OK, ${invalid} error(s).${missingText} Download started.`:"Excel URLs filled. Download started.");}
function setRiegelProgress(percent,label,active=false){const wrap=document.getElementById("riegelProgress"); const bar=document.getElementById("riegelProgressBar"); const text=document.getElementById("riegelProgressLabel"); if(!wrap||!bar||!text)return; wrap.style.display="block"; wrap.classList.toggle("active",active); bar.style.width=`${Math.max(0,Math.min(100,percent))}%`; text.textContent=label;}
function riegelHits(){return currentRows.filter(row=>row.match===true);}
function updateRiegelCounts(){const scanned=document.getElementById("riegelScannedCount"); const found=document.getElementById("riegelFoundCount"); const errors=document.getElementById("riegelErrorCount"); if(scanned)scanned.textContent=String(currentRows.length); if(found)found.textContent=String(riegelHits().length); if(errors)errors.textContent=String(currentRows.filter(row=>["error","skipped"].includes(String(row.status||"").toLowerCase())).length);}
async function parseJsonResponse(res){const text=await res.text(); try{return JSON.parse(text||"{}");}catch(err){throw new Error(text.trim().startsWith("<")?"Server returned HTML instead of JSON. Try again after refresh.":"Server returned invalid JSON.");}}
async function pollRiegelJob(jobId){while(true){const data=await api(`/api/job?id=${encodeURIComponent(jobId)}`); currentRows=data.rows||[]; renderRiegelTable(); updateRiegelCounts(); setRiegelProgress(data.percent||15,data.message||"Scanning images...",!data.done); status(`${data.message||"Scanning images..."} Scanned ${data.scanned||currentRows.length||0}. Riegel found ${data.found||riegelHits().length}. Issues/skips ${data.issues||0}.`); if(data.done){if(!data.ok)throw new Error(data.message||"Riegel scan failed"); setRiegelProgress(100,"Finished.",false); renderRiegelTable(); updateRiegelCounts(); return data;} await new Promise(resolve=>setTimeout(resolve,1000));}}
async function scanRiegelLogos(){const urls=document.getElementById("riegelUrls").value; const folder=document.getElementById("riegelFolder").value; const recursive=document.getElementById("riegelRecursive").checked; const btn=document.getElementById("riegelScanBtn"); if(!urls.trim()&&!folder.trim()){status("Paste image URLs or enter a folder path.");return;} if(btn)btn.disabled=true; currentRows=[]; renderRiegelTable(); updateRiegelCounts(); status("Starting Riegel scan..."); setRiegelProgress(8,"Starting scan...",true); try{const res=await fetch("/api/scan-riegel-logo",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify({urls,folder,recursive})}); const data=await parseJsonResponse(res); if(!res.ok)throw new Error(data.error||"Riegel scan failed"); if(data.job_id){await pollRiegelJob(data.job_id);return;} currentRows=data.rows||[]; renderRiegelTable(); updateRiegelCounts(); setRiegelProgress(100,"Finished.",false); status(`Scanned ${data.scanned} image(s). Riegel found in ${data.found}. Errors/skips: ${data.issues}.`);}finally{if(btn)btn.disabled=false;}}
function renderRiegelTable(){const wrap=document.getElementById("tableWrap"); if(!wrap)return; if(!currentRows.length){wrap.textContent="No scan results.";return;} const rows=currentRows.slice(0,1200).map(row=>{const hit=row.match===true; const color=hit?"#b42318":(row.status==="error"?"#b42318":"#0f766e"); const link=row.url?`<a href="${row.url}" target="_blank">${row.url}</a>`:""; const evidence=row.score?`Score ${row.score}${row.reference?`<div class="hint">${row.reference}</div>`:""}`:(row.text||""); return `<tr><td>${row.source||""}</td><td>${row.file||""}</td><td><strong style="color:${color}">${row.status||""}</strong><div class="hint">${row.message||""}</div></td><td>${row.path||""}</td><td>${link}</td><td>${evidence}</td></tr>`;}).join(""); const note=currentRows.length>1200?`<p class="hint">Showing first 1200 of ${currentRows.length} rows.</p>`:""; wrap.innerHTML=`${note}<table><thead><tr><th>Source</th><th>File</th><th>Status</th><th>Path</th><th>URL</th><th>Evidence</th></tr></thead><tbody>${rows}</tbody></table>`;}
async function copyRiegelHits(){const hits=riegelHits().map(row=>row.url||row.path).filter(Boolean); if(!hits.length){status("No Riegel matches to copy.");return;} await navigator.clipboard.writeText([...new Set(hits)].join("\\n")); status(`Copied ${hits.length} Riegel match path(s)/URL(s).`);}
function exportRiegelHits(){const hits=riegelHits(); if(!hits.length){status("No Riegel matches to export.");return;} const header=["source","file","status","path","url","message","score","reference","ocr_text"]; const lines=hits.map(row=>[row.source,row.file,row.status,row.path,row.url,row.message,row.score,row.reference,row.text].map(value=>`"${String(value||"").replace(/"/g,'""')}"`).join(",")); const blob=new Blob([[header.join(","),...lines].join("\\n")],{type:"text/csv"}); const a=document.createElement("a"); a.href=URL.createObjectURL(blob); a.download="riegel-logo-matches.csv"; a.click(); URL.revokeObjectURL(a.href); status(`Exported ${hits.length} Riegel match row(s).`);}
function renderUploadTable(){if(!currentRows.length){document.getElementById("tableWrap").textContent="No results.";return;} const rows=currentRows.slice(0,1000).map(row=>{const yearStyle=row.status==="year found"?' style="color:#b42318;font-weight:800"':""; return `<tr><td>${row.step||"upload"}</td><td>${row.file||row.product||""}</td><td${yearStyle}>${row.status||""}</td><td>${row.path||row.target||""}</td><td>${row.url?`<a href="${row.url}" target="_blank">${row.url}</a>`:(row.message||"")}</td></tr>`;}).join(""); document.getElementById("tableWrap").innerHTML=`<table><thead><tr><th>Step</th><th>File / Product</th><th>Status</th><th>Path</th><th>URL / Message</th></tr></thead><tbody>${rows}</tbody></table>`;}
function uploadLinks(){const finalRows=currentRows.filter(row=>row.step==="url"&&row.url); const source=finalRows.length?finalRows:currentRows.filter(row=>row.url); return [...new Set(source.map(row=>row.url))].join("\\n");}
async function copyUploadLinks(){const links=uploadLinks(); if(!links){status("No links to copy yet.");return;} await navigator.clipboard.writeText(links); status(`Copied ${links.split("\\n").length} links to clipboard.`);}
function downloadUploadLinks(){const links=uploadLinks(); if(!links){status("No links to export yet.");return;} const blob=new Blob([links],{type:"text/plain"}); const a=document.createElement("a"); a.href=URL.createObjectURL(blob); a.download=`primex-upload-links-${document.getElementById("uploadPath")?.value||"upload"}.txt`.replace(/[\\\\/]+/g,"-"); a.click(); URL.revokeObjectURL(a.href); status(`Exported ${links.split("\\n").length} links.`);}
document.getElementById("topFolder")?.addEventListener("change",loadProducts);
document.getElementById("topSearch")?.addEventListener("input",()=>{renderTopFolders(); loadProducts().catch(err=>status(err.message));});
document.getElementById("productSearch")?.addEventListener("input",()=>renderProducts());
document.getElementById("recentFolder")?.addEventListener("change",event=>{const value=event.target.value; if(!value)return; const top=document.getElementById("topFolder"); if(top){document.getElementById("topSearch").value=""; renderTopFolders(value); top.value=value; loadProducts().catch(err=>status(err.message));}});
document.getElementById("generateBtn")?.addEventListener("click",()=>generateUrls().catch(err=>status(err.message)));
document.getElementById("checkUrlBtn")?.addEventListener("click",()=>checkUrls().catch(err=>status(err.message)));
document.getElementById("gridBtn")?.addEventListener("click",()=>generateGrid().catch(err=>status(err.message)));
document.getElementById("organizeBtn")?.addEventListener("click",()=>organizePhotos().catch(err=>status(err.message)));
document.getElementById("uploadBtn")?.addEventListener("click",()=>uploadFiles().catch(err=>status(err.message)));
async function calculateGewicht(){
  const input=document.getElementById("gewichtFile");
  const file=input?.files[0];
  if(!file){status("Zgjidh një skedar Excel.");return;}
  const button=document.getElementById("gewichtBtn");
  button.disabled=true;
  status("Duke përpunuar Excel-in...");
  try{
    const prefix=document.getElementById("gewichtPrefix")?.value.trim();
    if(!/^[A-Za-z0-9]+$/.test(prefix||"")){throw new Error("Shkruaj një prefiks me shkronja ose numra.");}
    const form=new FormData();form.append("file",file);form.append("prefix",prefix);
    form.append("input_unit",document.getElementById("gewichtUnit")?.value||"g");
    const codeSuffix=document.getElementById("gewichtCodeSuffix")?.value.trim()||"";
    if(codeSuffix&&!/^[A-Za-z0-9]+$/.test(codeSuffix)){throw new Error("Prapashtesa duhet të ketë vetëm shkronja ose numra.");}
    form.append("code_suffix",codeSuffix);
    const response=await fetch("/api/gewicht",{method:"POST",body:form});
    if(!response.ok){const error=await response.json();throw new Error(error.error||"Përpunimi dështoi.");}
    const blob=await response.blob();
    const url=URL.createObjectURL(blob);
    const link=document.createElement("a");link.href=url;link.download="GEWICHT-"+file.name.replace(/\.(xlsx|xlsm)$/i,"")+".xlsx";
    document.body.appendChild(link);link.click();link.remove();
    setTimeout(()=>URL.revokeObjectURL(url),60000);
    status("U përpunuan "+response.headers.get("X-Gewicht-Rows")+" rreshta. Shkarkimi filloi.");
  }finally{button.disabled=false;}
}
document.getElementById("gewichtBtn")?.addEventListener("click",()=>calculateGewicht().catch(err=>status(err.message)));
document.getElementById("excelBtn")?.addEventListener("click",()=>fillExcelUrls().catch(err=>status(err.message)));
document.getElementById("riegelScanBtn")?.addEventListener("click",()=>scanRiegelLogos().catch(err=>status(err.message)));
document.getElementById("copyRiegelHitsBtn")?.addEventListener("click",()=>copyRiegelHits().catch(err=>status(err.message)));
document.getElementById("exportRiegelHitsBtn")?.addEventListener("click",exportRiegelHits);
document.getElementById("copyUploadLinksBtn")?.addEventListener("click",()=>copyUploadLinks().catch(err=>status(err.message)));
document.getElementById("downloadUploadLinksBtn")?.addEventListener("click",downloadUploadLinks);
document.getElementById("refreshBtn")?.addEventListener("click",()=>loadTopFolders().catch(err=>status(err.message)));
document.getElementById("uploadPath")?.addEventListener("input",()=>updateUploadPathSuggestions().catch(err=>status(err.message)));
document.getElementById("resultFilter")?.addEventListener("change",renderCurrentTable);
document.getElementById("showThumbs")?.addEventListener("change",renderCurrentTable);
document.getElementById("exportErrorsBtn")?.addEventListener("click",exportProblemRows);
document.getElementById("copyBtn")?.addEventListener("click",async()=>{await navigator.clipboard.writeText(document.getElementById("urlBox").value); status("Copied URLs to clipboard.");});
document.getElementById("downloadBtn")?.addEventListener("click",()=>{const blob=new Blob([document.getElementById("urlBox").value],{type:"text/plain"}); const a=document.createElement("a"); a.href=URL.createObjectURL(blob); a.download=`primex-urls-${document.getElementById("topFolder").value||"folder"}.txt`; a.click(); URL.revokeObjectURL(a.href);});
loadTopFolders().catch(err=>status(err.message));
</script>
"""


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt: str, *args: object) -> None:
        print(f"{self.address_string()} - {fmt % args}")

    def is_authenticated(self) -> bool:
        cookie = parse_cookies(self.headers.get("Cookie", "")).get("vs_session", "")
        return is_valid_session(cookie)

    def send_redirect(self, location: str, extra_headers: dict[str, str] | None = None) -> None:
        self.send_response(302)
        self.send_header("Location", location)
        for name, value in (extra_headers or {}).items():
            self.send_header(name, value)
        self.end_headers()

    def send_login(self, error: str = "") -> None:
        body = LOGIN_PAGE.replace("__ERROR_DISPLAY__", "block" if error else "none").replace("__ERROR__", html.escape(error)).encode("utf-8")
        self.send_html(body)

    def send_json(self, payload: object, status: int = 200) -> None:
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_html(self, body: bytes) -> None:
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_download(self, path: Path, filename: str, content_type: str, extra_headers: dict[str, str] | None = None) -> None:
        data = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        for name, value in (extra_headers or {}).items():
            self.send_header(name, value)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def send_asset(self, filename: str) -> bool:
        safe_name = Path(filename).name
        if safe_name not in {"logo-vs.png", "logoprimexeu.png"}:
            return False
        path = ASSET_ROOT / safe_name
        if not path.exists() or not path.is_file():
            return False
        data = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", "image/png")
        self.send_header("Cache-Control", "public, max-age=86400")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)
        return True

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        route = parsed.path.rstrip("/") or "/"
        try:
            if route.startswith("/assets/"):
                if not self.send_asset(route.rsplit("/", 1)[-1]):
                    self.send_error(404)
                return
            if route == "/login":
                if self.is_authenticated():
                    self.send_redirect("/")
                else:
                    self.send_login()
                return
            if route == "/logout":
                self.send_redirect("/login", {"Set-Cookie": "vs_session=; Path=/; HttpOnly; SameSite=Lax; Max-Age=0"})
                return
            if not self.is_authenticated():
                if route.startswith("/api/"):
                    self.send_json({"error": "Login required."}, 401)
                else:
                    self.send_redirect("/login")
                return
            if route == "/":
                self.send_html(shell("Upload Images", "upload", UPLOAD_PAGE))
                return
            if route == "/url-generator":
                self.send_html(shell("Primex URL Generator", "urls", URL_PAGE))
                return
            if route == "/image-generator":
                self.send_html(shell("Image Grid Generator", "grid", GRID_PAGE))
                return
            if route == "/photo-organizer":
                self.send_html(shell("Photo Organizer", "organizer", ORGANIZER_PAGE))
                return
            if route == "/upload":
                self.send_html(shell("Upload Images", "upload", UPLOAD_PAGE))
                return
            if route == "/excel-url-filler":
                self.send_html(shell("Excel URL Filler", "excel", EXCEL_PAGE))
                return
            if route == "/gewicht":
                self.send_html(shell("GEWICHT", "gewicht", GEWICHT_PAGE))
                return
            if route == "/riegel-logo-scanner":
                self.send_html(shell("Riegel Logo Scanner", "riegel", RIEGEL_PAGE))
                return
            if route == "/api/top-folders":
                self.send_json({"root": str(VS_ROOT), "baseUrl": BASE_URL, "folders": list_top_folders()})
                return
            if route == "/api/products":
                q = parse_qs(parsed.query)
                self.send_json({"products": list_product_folders(q.get("folder", [""])[0])})
                return
            if route == "/api/generate":
                q = parse_qs(parsed.query)
                rows = generate_urls(
                    q.get("folder", [""])[0],
                    q.get("product", [""])[0],
                    q.get("recursive", ["1"])[0] == "1",
                    q.get("cache", ["0"])[0] == "1",
                )
                self.send_json({"count": len(rows), "rows": rows})
                return
            if route == "/api/check-urls":
                q = parse_qs(parsed.query)
                rows = generate_urls(q.get("folder", [""])[0], q.get("product", [""])[0], q.get("recursive", ["1"])[0] == "1")
                online = 0
                for row in rows:
                    ok, message = check_public_url(row["url"])
                    row["ok"] = ok
                    row["status"] = message
                    if ok:
                        online += 1
                self.send_json({"count": len(rows), "online": online, "errors": len(rows) - online, "rows": rows})
                return
            if route == "/api/check-url":
                q = parse_qs(parsed.query)
                url = q.get("url", [""])[0].strip()
                if not url.startswith(BASE_URL):
                    raise ValueError("URL must use the configured Primex photos base.")
                ok, message = check_public_url(url)
                self.send_json({"url": url, "ok": ok, "status": message})
                return
            if route == "/api/job":
                q = parse_qs(parsed.query)
                job = job_snapshot(q.get("id", [""])[0])
                if not job:
                    self.send_json({"error": "Job not found"}, 404)
                else:
                    self.send_json(job)
                return
            if route == "/api/generate-images":
                q = parse_qs(parsed.query)
                rows = generate_grid_images(
                    q.get("folder", [""])[0],
                    q.get("product", [""])[0],
                    [item for item in q.get("quantities", [""])[0].split(",") if item],
                    q.get("format", ["jpg"])[0],
                    q.get("delete", ["0"])[0] == "1",
                    q.get("cache", ["0"])[0] == "1",
                )
                created = sum(1 for row in rows if row.get("status") == "created")
                self.send_json({"rows": rows, "created": created, "issues": len(rows) - created})
                return
            if route == "/api/organize-photos":
                q = parse_qs(parsed.query)
                rows = organize_photos(
                    q.get("source", [""])[0],
                    q.get("destination", [""])[0],
                    q.get("move", ["0"])[0] == "1",
                )
                changed = sum(1 for row in rows if row.get("status") in {"copied", "moved"})
                self.send_json({"rows": rows, "changed": changed, "issues": len(rows) - changed})
                return
            self.send_json({"error": "Not found"}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 400)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        route = parsed.path.rstrip("/") or "/"
        try:
            if route == "/login":
                length = int(self.headers.get("Content-Length", "0") or "0")
                fields = parse_qs(self.rfile.read(length).decode("utf-8", errors="replace"))
                email = fields.get("email", [""])[0].strip()
                password = fields.get("password", [""])[0]
                if hmac.compare_digest(email, LOGIN_EMAIL) and hmac.compare_digest(password, LOGIN_PASSWORD):
                    self.send_redirect("/", {"Set-Cookie": f"vs_session={make_session_cookie()}; Path=/; HttpOnly; SameSite=Lax; Max-Age=86400"})
                else:
                    self.send_login("Wrong email or password.")
                return
            if not self.is_authenticated():
                self.send_json({"error": "Login required."}, 401)
                return
            if route == "/api/scan-riegel-logo":
                length = int(self.headers.get("Content-Length", "0") or "0")
                try:
                    payload = json.loads(self.rfile.read(length).decode("utf-8", errors="replace") or "{}")
                except json.JSONDecodeError:
                    raise ValueError("Expected JSON body.")
                url_text = str(payload.get("urls", ""))
                folder_text = str(payload.get("folder", ""))
                if not url_text.strip() and not folder_text.strip():
                    raise ValueError("Paste image URLs or enter a folder path to scan.")
                job_id = secrets.token_urlsafe(12)
                with JOBS_LOCK:
                    JOBS[job_id] = {
                        "done": False,
                        "ok": True,
                        "phase": "queued",
                        "message": "Riegel scan queued...",
                        "percent": 5,
                        "scanned": 0,
                        "found": 0,
                        "issues": 0,
                        "rows": [],
                        "started": time.time(),
                    }
                thread = threading.Thread(
                    target=run_riegel_scan_job,
                    args=(job_id, url_text, folder_text, bool(payload.get("recursive", True))),
                    daemon=True,
                )
                thread.start()
                self.send_json({"job_id": job_id, "done": False})
                return
            if route not in {"/api/upload", "/api/fill-excel", "/api/gewicht"}:
                self.send_json({"error": "Not found"}, 404)
                return
            content_type = self.headers.get("Content-Type", "")
            if not content_type.startswith("multipart/form-data"):
                raise ValueError("Expected multipart upload.")
            form = cgi.FieldStorage(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": content_type,
                    "CONTENT_LENGTH": self.headers.get("Content-Length", "0"),
                },
            )
            if route == "/api/gewicht":
                field = form["file"] if "file" in form else None
                if field is None or not getattr(field, "filename", ""):
                    raise ValueError("Choose an Excel file.")
                original_name = safe_upload_filename(field.filename)
                suffix = Path(original_name).suffix.lower()
                if suffix not in {".xlsx", ".xlsm"}:
                    raise ValueError("Upload an .xlsx or .xlsm file.")
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp:
                    source_path = Path(temp.name)
                    while chunk := field.file.read(1024 * 1024):
                        temp.write(chunk)
                output_path = None
                try:
                    output_path, count = build_gewicht_workbook(source_path, form.getfirst("prefix", "AM"), form.getfirst("input_unit", "g"), form.getfirst("code_suffix", ""))
                    self.send_download(
                        output_path,
                        "GEWICHT.xlsx",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        {"X-Gewicht-Rows": str(count), "Access-Control-Expose-Headers": "X-Gewicht-Rows"},
                    )
                finally:
                    source_path.unlink(missing_ok=True)
                    if output_path is not None:
                        output_path.unlink(missing_ok=True)
                return
            if route == "/api/fill-excel":
                field = form["file"] if "file" in form else None
                if field is None or not getattr(field, "filename", ""):
                    raise ValueError("Choose an Excel file.")
                original_name = safe_upload_filename(field.filename)
                if Path(original_name).suffix.lower() not in {".xlsm", ".xlsx"}:
                    raise ValueError("Upload an .xlsm or .xlsx file.")
                with tempfile.NamedTemporaryFile(delete=False, suffix=Path(original_name).suffix) as temp:
                    while True:
                        chunk = field.file.read(1024 * 1024)
                        if not chunk:
                            break
                        temp.write(chunk)
                    source_path = Path(temp.name)
                output_path, stats = fill_excel_urls_zip_patch(
                    source_path,
                    original_name,
                    form.getfirst("top", "KR").strip() or "KR",
                    form.getfirst("prefix", "KR").strip(),
                    form.getfirst("suffix", "bio").strip(),
                    form.getfirst("cache_bust", "0") == "1",
                )
                self.send_download(
                    output_path,
                    output_path.name,
                    "application/vnd.ms-excel.sheet.macroEnabled.12",
                    {"X-Excel-Url-Stats": json.dumps(stats)},
                )
                try:
                    source_path.unlink(missing_ok=True)
                except Exception:
                    pass
                return
            destination = resolve_upload_path(form.getfirst("path", ""))
            overwrite = form.getfirst("overwrite", "0") == "1"
            check_years = form.getfirst("check_years", "0") == "1"
            automate = form.getfirst("automate", "0") == "1"
            export_format = form.getfirst("format", "jpg")
            delete_source = form.getfirst("delete", "0") == "1"
            cache_bust = form.getfirst("cache_bust", "0") == "1"
            automate_only = form.getfirst("automate_only", "0") == "1"
            reported_uploaded = max(0, int(form.getfirst("uploaded_count", "0") or "0"))
            quantities = [item for item in form.getfirst("quantities", "1,3,6,12").split(",") if item]
            relative_parts = list(destination.relative_to(VS_ROOT).parts)
            destination.mkdir(parents=True, exist_ok=True)
            fields = form["files"] if "files" in form else []
            if not isinstance(fields, list):
                fields = [fields]
            rows = []
            for field in fields:
                if not getattr(field, "filename", ""):
                    continue
                try:
                    filename = safe_upload_filename(field.filename)
                    if Path(filename).suffix.lower() not in URL_EXTENSIONS:
                        rows.append({"step": "upload", "file": filename, "status": "skipped", "message": "File type is not allowed"})
                        continue
                    target = destination / filename
                    if target.exists() and not overwrite:
                        rows.append({"step": "upload", "file": filename, "status": "skipped", "path": str(target.relative_to(VS_ROOT)).replace(os.sep, "/"), "message": "File exists"})
                        continue
                    with target.open("wb") as out:
                        while True:
                            chunk = field.file.read(1024 * 1024)
                            if not chunk:
                                break
                            out.write(chunk)
                    rows.append({
                        "step": "upload",
                        "file": filename,
                        "status": "uploaded",
                        "path": str(target.relative_to(VS_ROOT)).replace(os.sep, "/"),
                        "url": public_url_for(target, cache_bust),
                    })
                    if check_years:
                        year_result = detect_years_in_image(target)
                        if year_result["checked"]:
                            years = year_result.get("years", [])
                            rows.append({
                                "step": "year check",
                                "file": filename,
                                "status": "year found" if years else "no year",
                                "path": str(target.relative_to(VS_ROOT)).replace(os.sep, "/"),
                                "message": year_result["message"],
                            })
                        else:
                            rows.append({
                                "step": "year check",
                                "file": filename,
                                "status": "skipped",
                                "path": str(target.relative_to(VS_ROOT)).replace(os.sep, "/"),
                                "message": year_result["message"],
                            })
                except Exception as exc:
                    rows.append({"step": "upload", "file": getattr(field, "filename", ""), "status": "error", "message": str(exc)})
            uploaded = sum(1 for row in rows if row.get("status") == "uploaded")
            years_found = sum(1 for row in rows if row.get("status") == "year found")
            created = 0
            url_count = 0
            if automate and (uploaded or automate_only):
                job_uploaded = uploaded or reported_uploaded
                if delete_source:
                    for row in rows:
                        if row.get("step") == "upload" and row.get("status") == "uploaded":
                            row.pop("url", None)
                            row["message"] = "Temporary source; removed after successful grid generation"
                job_id = secrets.token_urlsafe(12)
                with JOBS_LOCK:
                    JOBS[job_id] = {
                        "done": False,
                        "ok": True,
                        "phase": "queued",
                        "message": "Upload finished. Starting automation...",
                        "percent": 8,
                        "uploaded": job_uploaded,
                        "created": 0,
                        "urls": 0,
                        "years": years_found,
                        "issues": 0,
                        "rows": rows,
                        "started": time.time(),
                    }
                thread = threading.Thread(
                    target=run_upload_automation_job,
                    args=(job_id, relative_parts, quantities, export_format, delete_source, cache_bust),
                    daemon=True,
                )
                thread.start()
                self.send_json({"job_id": job_id, "done": False, "rows": rows, "uploaded": job_uploaded, "created": 0, "urls": 0, "years": years_found, "issues": 0})
                return

            successful = sum(1 for row in rows if row.get("status") in {"uploaded", "uploaded then organized", "moved", "copied", "created", "ready", "year found", "no year"})
            self.send_json({"rows": rows, "uploaded": uploaded, "created": created, "urls": url_count, "years": years_found, "issues": len(rows) - successful})
        except Exception as exc:
            self.send_json({"error": str(exc)}, 400)


def main() -> None:
    if not VS_ROOT.exists():
        raise SystemExit(f"VS_ROOT does not exist: {VS_ROOT}")
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"VS Tools serving on http://{HOST}:{PORT}")
    print(f"Scanning {VS_ROOT}")
    server.serve_forever()


if __name__ == "__main__":
    main()
