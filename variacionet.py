from __future__ import annotations

import csv
import re
from openpyxl import load_workbook


ACCOUNT_IDS = {"visando": 24765, "versanel": 8942}


def account_id_for(account):
    try:
        return ACCOUNT_IDS[account.lower()]
    except (AttributeError, KeyError):
        raise ValueError("Zgjidh Visando ose Versanel para perpunimit.") from None


def ensure_xlsx_columns(ws):
    columns = {normalize_header(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}
    for name in ("amazon_mengen_abgleich", "amazon_mengen_abgleich_account_id", "amazon_mengen_abgleich_max"):
        if name not in columns:
            col = ws.max_column + 1
            ws.cell(1, col).value = name
            columns[name] = col
    return columns


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(value):
    return re.sub(r"[^a-z0-9]+", "_", clean_text(value).lower()).strip("_")


def parse_sku(sku):
    """
    Returns (base_sku, quantity) or (None, None).

    Supported examples:
      01-RO20001, 03-RO20001, 0006-RO20001
      RO20001-01, RO20001-03, RO20001-0006
    """
    sku = clean_text(sku)
    if not sku or "-" not in sku:
        return None, None

    # Quantity at the beginning: 03-RO20001
    match = re.match(r"^(\d+)-(.+)$", sku)
    if match:
        qty = int(match.group(1))
        base = match.group(2).strip()
        if qty > 0 and base:
            return base.upper(), qty

    # Quantity at the end: RO20001-03
    match = re.match(r"^(.+)-(\d+)$", sku)
    if match:
        base = match.group(1).strip()
        qty = int(match.group(2))
        if qty > 0 and base:
            return base.upper(), qty

    return None, None
def find_columns(headers):
    normalized = {normalize_header(value): index for index, value in enumerate(headers, start=1)}

    sku_aliases = ["druck_pseudonym", "sku", "druckpseudonym"]
    lager_aliases = ["lager_nr", "lagernr", "lager_number"]
    set_aliases = ["set_article", "setartikel", "set_artikel"]

    def first_match(aliases, fallback):
        for alias in aliases:
            if alias in normalized:
                return normalized[alias]
        return fallback

    # Falls back to A, B, C if headers are not recognized.
    return (
        first_match(sku_aliases, 1),
        first_match(lager_aliases, 2),
        first_match(set_aliases, 3),
    )


def process_xlsx(input_path, output_path, sheet_name=None, account=None):
    account_id = account_id_for(account)
    keep_vba = input_path.lower().endswith(".xlsm")
    wb = load_workbook(input_path, keep_vba=keep_vba)

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    headers = [ws.cell(1, col).value for col in range(1, max(ws.max_column, 3) + 1)]
    sku_col, lager_col, set_col = find_columns(headers)
    amazon_cols = ensure_xlsx_columns(ws)

    parents = {}
    duplicate_parents = []
    parsed_rows = {}

    # First pass: find quantity-1 parent anywhere in the sheet.
    for row in range(2, ws.max_row + 1):
        sku = clean_text(ws.cell(row, sku_col).value)
        lager = clean_text(ws.cell(row, lager_col).value)
        base, qty = parse_sku(sku)
        parsed_rows[row] = (base, qty)

        if base and qty == 1 and lager:
            if base not in parents:
                parents[base] = lager
            elif parents[base] != lager:
                duplicate_parents.append(
                    f"{base}: LagerNr {parents[base]} dhe {lager}"
                )

    filled = 0
    already_filled = 0
    parent_rows = 0
    no_parent = []
    invalid_sku = []

    # Second pass: fill Amazon fields and only blank set_article cells.
    for row in range(2, ws.max_row + 1):
        if all(ws.cell(row, col).value is None for col in range(1, ws.max_column + 1)):
            continue
        ws.cell(row, amazon_cols["amazon_mengen_abgleich"]).value = 1
        ws.cell(row, amazon_cols["amazon_mengen_abgleich_account_id"]).value = account_id
        max_cell = ws.cell(row, amazon_cols["amazon_mengen_abgleich_max"])
        if not clean_text(max_cell.value):
            max_cell.value = 12 if parsed_rows[row][1] == 1 else 4
        current_set = clean_text(ws.cell(row, set_col).value)
        if current_set:
            already_filled += 1
            continue

        sku = clean_text(ws.cell(row, sku_col).value)
        base, qty = parsed_rows[row]

        if not sku:
            continue

        if not base or qty is None:
            invalid_sku.append(f"Rreshti {row}: {sku}")
            continue

        if qty == 1:
            parent_rows += 1
            continue

        parent_lager = parents.get(base)
        if not parent_lager:
            no_parent.append(f"Rreshti {row}: {sku}")
            continue

        cell = ws.cell(row, set_col)
        cell.number_format = "@"
        cell.value = f"{parent_lager}:{qty}"
        filled += 1

    wb.save(output_path)

    return {
        "sheet": ws.title,
        "filled": filled,
        "already_filled": already_filled,
        "parent_rows": parent_rows,
        "no_parent": no_parent,
        "invalid_sku": invalid_sku,
        "duplicate_parents": duplicate_parents,
    }


def detect_csv_encoding(path):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                f.read()
            return enc
        except UnicodeDecodeError:
            pass
    return "latin-1"


def detect_csv_dialect(path, encoding):
    with open(path, "r", encoding=encoding, newline="") as f:
        sample = f.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;	|")
    except csv.Error:
        return csv.excel


def process_csv(input_path, output_path, account=None):
    account_id = account_id_for(account)
    encoding = detect_csv_encoding(input_path)
    dialect = detect_csv_dialect(input_path, encoding)

    with open(input_path, "r", encoding=encoding, newline="") as f:
        rows = list(csv.reader(f, dialect))

    if not rows:
        raise ValueError("Fajlli CSV është bosh.")

    headers = rows[0]
    while len(headers) < 3:
        headers.append("")

    sku_col, lager_col, set_col = find_columns(headers)
    sku_idx, lager_idx, set_idx = sku_col - 1, lager_col - 1, set_col - 1
    columns = {normalize_header(value): index for index, value in enumerate(headers)}
    for name in ("amazon_mengen_abgleich", "amazon_mengen_abgleich_account_id", "amazon_mengen_abgleich_max"):
        if name not in columns:
            columns[name] = len(headers)
            headers.append(name)

    max_idx = max(sku_idx, lager_idx, set_idx, *columns.values())
    for row in rows:
        while len(row) <= max_idx:
            row.append("")

    parents = {}
    duplicate_parents = []
    parsed_rows = {}

    for index, row in enumerate(rows[1:], start=2):
        sku = clean_text(row[sku_idx])
        lager = clean_text(row[lager_idx])
        base, qty = parse_sku(sku)
        parsed_rows[index] = (base, qty)

        if base and qty == 1 and lager:
            if base not in parents:
                parents[base] = lager
            elif parents[base] != lager:
                duplicate_parents.append(
                    f"{base}: LagerNr {parents[base]} dhe {lager}"
                )

    filled = 0
    already_filled = 0
    parent_rows = 0
    no_parent = []
    invalid_sku = []

    for index, row in enumerate(rows[1:], start=2):
        if not any(clean_text(value) for value in row):
            continue
        row[columns["amazon_mengen_abgleich"]] = "1"
        row[columns["amazon_mengen_abgleich_account_id"]] = str(account_id)
        max_column = columns["amazon_mengen_abgleich_max"]
        if not clean_text(row[max_column]):
            row[max_column] = "12" if parsed_rows[index][1] == 1 else "4"

        if clean_text(row[set_idx]):
            already_filled += 1
            continue

        sku = clean_text(row[sku_idx])
        base, qty = parsed_rows[index]

        # Plotëso kolonën H
        if not sku:
            continue

        if not base or qty is None:
            invalid_sku.append(f"Rreshti {index}: {sku}")
            continue

        if qty == 1:
            parent_rows += 1
            continue

        parent_lager = parents.get(base)
        if not parent_lager:
            no_parent.append(f"Rreshti {index}: {sku}")
            continue

        row[set_idx] = f"{parent_lager}:{qty}"
        filled += 1

    with open(output_path, "w", encoding=encoding, newline="") as f:
        writer = csv.writer(
            f,
            delimiter=dialect.delimiter,
            quoting=csv.QUOTE_ALL
        )
        writer.writerows(rows)

    return {
        "sheet": "CSV",
        "filled": filled,
        "already_filled": already_filled,
        "parent_rows": parent_rows,
        "no_parent": no_parent,
        "invalid_sku": invalid_sku,
        "duplicate_parents": duplicate_parents,
    }


