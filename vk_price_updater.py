"""Workbook logic and web page for the VK Price Updater."""

from __future__ import annotations

import re
from copy import copy
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TARGET_HEADER = "VK PARENT & CHILDREN"
OUTPUT_HEADER = "VK AFTER DISCOUNTS"
NUMBER_TEXT = re.compile(r"^([+-]?)(\d+(?:\.\d*)?|\.\d+)$")
PRODUCT_PREFIX = re.compile(r"^\s*(01|03|06|12)(?=\D|$)")
PREFIX_DISCOUNTS = {"01": Decimal("0.5"), "03": Decimal("1"), "06": Decimal("1"), "12": Decimal("2")}
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM_MIME = "application/vnd.ms-excel.sheet.macroEnabled.12"


@dataclass
class ScanResult:
    matches: int
    changes: int
    sheets: list[str]
    samples: list[dict[str, object]]
    pack_changes: int = 0

    def as_dict(self) -> dict[str, object]:
        return asdict(self)


def matching_columns(sheet) -> list[tuple[int, int]]:
    found = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)):
        for cell in row:
            if isinstance(cell.value, str) and " ".join(cell.value.split()).casefold() == TARGET_HEADER.casefold():
                found.append((cell.row, cell.column))
    return found


def adjusted_value(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            amount = Decimal(str(value))
        except InvalidOperation:
            return None
        if not amount.is_finite() or amount >= 50:
            return None
        changed = amount - 3
        return int(changed) if changed == changed.to_integral_value() else float(changed)
    if isinstance(value, str):
        stripped = value.strip()
        if not NUMBER_TEXT.fullmatch(stripped):
            return None
        try:
            amount = Decimal(stripped)
        except InvalidOperation:
            return None
        if amount >= 50:
            return None
        changed = amount - 3
        decimals = len(stripped.split(".", 1)[1]) if "." in stripped else 0
        rendered = f"{changed:.{decimals}f}" if decimals else str(changed)
        return value[: len(value) - len(value.lstrip())] + rendered + value[len(value.rstrip()) :]
    return None


def product_discount(product: object) -> Decimal | None:
    if not isinstance(product, str):
        return None
    match = PRODUCT_PREFIX.match(product)
    return PREFIX_DISCOUNTS[match.group(1)] if match else None


def discounted_value(value: object, discount: Decimal):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            amount = Decimal(str(value))
        except InvalidOperation:
            return None
        if not amount.is_finite():
            return None
        changed = amount - discount
        return int(changed) if changed == changed.to_integral_value() else float(changed)
    if isinstance(value, str):
        stripped = value.strip()
        if not NUMBER_TEXT.fullmatch(stripped):
            return None
        try:
            amount = Decimal(stripped)
        except InvalidOperation:
            return None
        if not amount.is_finite():
            return None
        changed = amount - discount
        original_decimals = len(stripped.split(".", 1)[1]) if "." in stripped else 0
        discount_decimals = max(0, -discount.as_tuple().exponent)
        decimals = max(original_decimals, discount_decimals)
        rendered = f"{changed:.{decimals}f}" if decimals else str(changed)
        return value[: len(value) - len(value.lstrip())] + rendered + value[len(value.rstrip()) :]
    return None


def scan_workbook(data: bytes, filename: str, *, apply: bool = False) -> tuple[ScanResult, bytes | None]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Choose an .xlsx or .xlsm workbook.")
    if not data:
        raise ValueError("The workbook is empty.")
    if len(data) > MAX_WORKBOOK_BYTES:
        raise ValueError("The workbook is larger than 25 MB.")

    workbook = load_workbook(BytesIO(data), keep_vba=suffix == ".xlsm")
    result = ScanResult(matches=0, changes=0, sheets=[], samples=[])
    try:
        for sheet in workbook.worksheets:
            headings = matching_columns(sheet)
            if not headings:
                continue
            result.sheets.append(sheet.title)
            result.matches += len(headings)
            for header_row, column in headings:
                output_header = OUTPUT_HEADER if len(headings) == 1 else f"{OUTPUT_HEADER} ({get_column_letter(column)})"
                existing_output_column = next(
                    (cell.column for cell in sheet[header_row] if isinstance(cell.value, str) and cell.value == output_header),
                    None,
                )
                output_column = existing_output_column
                if apply:
                    if output_column is None:
                        output_column = sheet.max_column + 1
                    source_header = sheet.cell(header_row, column)
                    target_header = sheet.cell(header_row, output_column)
                    target_header.value = output_header
                    target_header._style = copy(source_header._style)
                    source_width = sheet.column_dimensions[get_column_letter(column)].width
                    sheet.column_dimensions[get_column_letter(output_column)].width = max(source_width or 13, 22)
                for row in range(header_row + 1, sheet.max_row + 1):
                    cell = sheet.cell(row, column)
                    first_changed = None
                    if cell.data_type != "f" and existing_output_column is None:
                        first_changed = adjusted_value(cell.value)
                    first_price = cell.value if first_changed is None else first_changed
                    discounted = None
                    if cell.data_type != "f":
                        discount = product_discount(sheet.cell(row, 1).value)
                        if discount is not None:
                            discounted = discounted_value(first_price, discount)
                            if discounted is not None:
                                result.pack_changes += 1
                    final_price = first_price if discounted is None else discounted
                    if first_changed is not None or discounted is not None:
                        result.changes += 1
                        if len(result.samples) < 100:
                            result.samples.append({"sheet": sheet.title, "row": row, "before": str(first_price), "after": str(final_price)})
                    if apply:
                        cell.value = first_price
                        target_cell = sheet.cell(row, output_column)
                        target_cell.value = final_price
                        target_cell._style = copy(cell._style)
        if not apply:
            return result, None
        if not result.matches:
            raise ValueError(f'Column "{TARGET_HEADER}" was not found in the first 30 rows of any sheet.')
        output = BytesIO()
        workbook.save(output)
        return result, output.getvalue()
    finally:
        workbook.close()


PAGE = r"""
<h1>VK Price Updater</h1>
<p class="hint">Keep the VK price after the first calculation in <strong>VK PARENT &amp; CHILDREN</strong> and add <strong>VK AFTER DISCOUNTS</strong> for the Produkt code adjustment. Your uploaded file stays unchanged.</p>
<section class="panel">
  <h2>Select workbook</h2>
  <p class="hint">Choose an .xlsx or .xlsm file, up to 25 MB. The updater checks every sheet for the heading in its first 30 rows.</p>
  <div class="field"><label for="vkPriceFile">Excel workbook</label><input id="vkPriceFile" type="file" accept=".xlsx,.xlsm"></div>
  <div class="actions"><button type="button" id="vkPreviewBtn" class="secondary">Preview changes</button><button type="button" id="vkUpdateBtn">Update workbook</button></div>
  <div class="status" id="vkPriceStatus" role="status" aria-live="polite"></div>
</section>
<section class="panel">
  <h2>Preview</h2>
  <p class="hint" id="vkPriceSummary">Choose a workbook and click Preview changes.</p>
  <div id="vkPriceSamples"></div>
  <p class="hint">The VK PARENT &amp; CHILDREN column first reduces values below 50 by 3. VK AFTER DISCOUNTS then applies the Produkt code adjustment: 01 = 0.50 EUR, 03 or 06 = 1.00 EUR, 12 = 2.00 EUR. Formulas, empty cells, and other text are copied unchanged. Select the original workbook when updating.</p>
</section>
<script>
(() => {
  const fileInput = document.getElementById('vkPriceFile');
  const previewBtn = document.getElementById('vkPreviewBtn');
  const updateBtn = document.getElementById('vkUpdateBtn');
  const status = document.getElementById('vkPriceStatus');
  const summary = document.getElementById('vkPriceSummary');
  const samples = document.getElementById('vkPriceSamples');
  const escapeHtml = value => String(value).replace(/[&<>"']/g, char => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[char]));
  function message(value, error = false) {
    status.style.display = 'block';
    status.style.borderLeftColor = error ? '#dc2626' : '';
    status.textContent = value;
  }
  function selectedFile() {
    const file = fileInput.files?.[0];
    if (!file || !/\.(xlsx|xlsm)$/i.test(file.name)) throw new Error('Choose an .xlsx or .xlsm workbook.');
    if (file.size > 25 * 1024 * 1024) throw new Error('The workbook is larger than 25 MB.');
    return file;
  }
  function showPreview(data) {
    summary.textContent = data.matches
      ? `Found ${data.matches} matching column(s) on ${data.sheets.length} sheet(s). ${data.changes} value(s) will change; ${data.pack_changes || 0} use the Produkt code adjustment.`
      : 'No VK PARENT & CHILDREN heading was found in the first 30 rows of any sheet.';
    const rows = data.samples || [];
    samples.innerHTML = rows.length
      ? '<div style="overflow:auto"><table><thead><tr><th>Sheet</th><th>Row</th><th>Before</th><th>After</th></tr></thead><tbody>' + rows.map(row => `<tr><td>${escapeHtml(row.sheet)}</td><td>${escapeHtml(row.row)}</td><td>${escapeHtml(row.before)}</td><td>${escapeHtml(row.after)}</td></tr>`).join('') + '</tbody></table></div>'
      : '';
  }
  async function run(update) {
    const file = selectedFile();
    previewBtn.disabled = updateBtn.disabled = true;
    message(update ? 'Updating workbook...' : 'Scanning workbook...');
    try {
      const form = new FormData();
      form.append('file', file);
      const response = await fetch(update ? '/api/vk-prices/update' : '/api/vk-prices/preview', {method:'POST', body:form});
      if (!response.ok) {
        const result = await response.json().catch(() => ({}));
        throw new Error(result.error || 'Could not process the workbook.');
      }
      if (update) {
        const blob = await response.blob();
        const url = URL.createObjectURL(blob);
        const link = document.createElement('a');
        link.href = url;
        link.download = file.name.replace(/\.(xlsx|xlsm)$/i, '') + '_updated.' + file.name.split('.').pop().toLowerCase();
        document.body.appendChild(link);
        link.click();
        link.remove();
        setTimeout(() => URL.revokeObjectURL(url), 60000);
        message(`Updated ${response.headers.get('X-VK-Changes') || 'the'} value(s). Download started.`);
      } else {
        showPreview(await response.json());
        message('Preview ready. The original workbook has not been changed.');
      }
    } catch (error) {
      message(error.message, true);
    } finally {
      previewBtn.disabled = updateBtn.disabled = false;
    }
  }
  fileInput.addEventListener('change', () => { summary.textContent = 'Click Preview changes to scan the selected workbook.'; samples.replaceChildren(); status.style.display = 'none'; });
  previewBtn.addEventListener('click', () => run(false));
  updateBtn.addEventListener('click', () => run(true));
})();
</script>
"""
