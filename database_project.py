"""VS Database Project tools copied from the original Laravel feature.

The two database reports read imported table exports from private local files.
No connection to the original SQL Server is made by this module.
"""

from __future__ import annotations

import csv
import io
import math
import os
import re
import tempfile
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATA_DIR = Path(os.environ.get(
    "VS_DATABASE_DATA_DIR",
    Path(os.environ.get("XDG_DATA_HOME", Path.home() / ".local" / "share")) / "vs-database-project",
))
TABLE_FILES = {
    "active": "active_products.xlsx",
    "dreamrobot": "dreamrobot_products.xlsx",
}
PACKS = (
    ("12er", "12", 12, Decimal("4.33"), Decimal("2.84")),
    ("6er", "06", 6, Decimal("4.33"), Decimal("1.55")),
    ("3er", "03", 3, Decimal("4.33"), Decimal("0.85")),
    ("4er", "04", 4, Decimal("4.33"), Decimal("1.55")),
    ("2er", "02", 2, Decimal("3.60"), Decimal("0.68")),
    ("1er", "01", 1, Decimal("3.60"), Decimal("0.42")),
)
BATCH_LABELS = ("1er", "3er", "6er", "12er")


def money(value: object) -> Decimal:
    return Decimal(str(value).strip().replace(",", ".")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def factor(label: str, pricing_type: str, spalte: str = "W") -> Decimal:
    feinkost = pricing_type == "feinkost" or (pricing_type == "normal" and spalte.strip().upper() == "F")
    if feinkost:
        return Decimal("1.53") if label in {"12er", "6er", "3er"} else Decimal("1.48")
    return Decimal("1.56") if label in {"12er", "6er"} else Decimal("1.55")


def ending_99(raw: Decimal) -> Decimal:
    whole = math.floor(raw)
    return Decimal(whole) + (Decimal("0.99") if raw - whole >= Decimal("0.5") else Decimal("-0.01"))


def calculate_pack(ek1er: Decimal, label: str, pricing_type: str, spalte: str = "W", database: bool = False) -> dict[str, object]:
    pack = next(item for item in PACKS if item[0] == label)
    _, prefix, qty, porto, pick_pack = pack
    if database and label == "1er":
        pick_pack = Decimal("0.68")  # The original database report uses a different 1er pack cost.
    ekxer = (ek1er * qty).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    multiplier = factor(label, pricing_type, spalte)
    raw = (ekxer + porto + pick_pack) * multiplier
    return {"label": label, "prefix": prefix, "ekxer": ekxer, "porto": porto,
            "pick_pack": pick_pack, "factor": multiplier, "raw": raw, "vk": ending_99(raw)}


def _csv_rows(data: bytes, delimiter: str | None = None) -> list[list[object]]:
    try:
        content = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = data.decode("cp1252")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(content[:4096], delimiters=";,\t").delimiter
        except csv.Error:
            delimiter = ","
    return list(csv.reader(io.StringIO(content), delimiter=delimiter))


def read_rows(data: bytes, filename: str, *, sku_csv: bool = False) -> list[list[object]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".txt"}:
        return _csv_rows(data, ";" if sku_csv else None)
    if suffix in {".xlsx", ".xlsm"}:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return [list(row) for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError(".xls support requires xlrd on the server.") from exc
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    raise ValueError("Upload an .xlsx, .xlsm, .xls or .csv file.")


def _output_path() -> Path:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        return Path(tmp.name)


def _save_book(book: Workbook) -> Path:
    path = _output_path()
    book.save(path)
    return path


def _autosize(sheet, count: int) -> None:
    from openpyxl.utils import get_column_letter
    for index in range(1, count + 1):
        longest = max((len(str(sheet.cell(row, index).value or "")) for row in range(1, sheet.max_row + 1)), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 70)


def pricing_report(data: bytes, filename: str, pricing_type: str) -> tuple[Path, int]:
    if pricing_type not in {"normal", "feinkost"}:
        raise ValueError("Choose Normal or Feinkost pricing.")
    rows = read_rows(data, filename)
    book = Workbook()
    sheet = book.active
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "Pricing Calculation Type: " + pricing_type.upper()
    sheet["A1"].font = __import__("openpyxl").styles.Font(bold=True)
    sheet.append(["Produkt", "Ek1er prices", "EK PARENT & CHILDREN", "VK PARENT & CHILDREN", "Pack Label"])
    count = 0
    for row in rows[1:]:
        product = str(row[0]).strip() if row and row[0] is not None else ""
        try:
            ek1er = money(row[1])
        except (IndexError, ValueError, ArithmeticError):
            continue
        if not product or ek1er <= 0:
            continue
        spalte = str(row[2]).strip().upper() if len(row) > 2 and row[2] is not None else "W"
        sheet.append([product, f"{ek1er:.2f}", None, None, "Parent"])
        for label in BATCH_LABELS:
            result = calculate_pack(ek1er, label, pricing_type, spalte)
            sheet.append([f"{result['prefix']}-WW{product}", None, f"{result['ekxer']:.2f}", f"{result['vk']:.2f}", label])
        count += 1
    _autosize(sheet, 5)
    return _save_book(book), count


def sku_report(data: bytes, filename: str) -> tuple[Path, int]:
    rows = read_rows(data, filename, sku_csv=True)
    headers = ["a_seller_sku", "Article number without letter", "Article number with a line and two numbers in the end"]
    if rows and str(rows[0][0] or "").lstrip("\ufeff").strip().lower() == "a_seller_sku":
        if Path(filename).suffix.lower() in {".csv", ".txt"}:
            headers = [(str(rows[0][i]).lstrip("\ufeff") if len(rows[0]) > i and rows[0][i] else headers[i]) for i in range(3)]
        rows = rows[1:]
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        raw = row[0] if row else None
        sku = str(raw).strip() if raw is not None else ""
        sheet.append([raw, sku if sku and not re.search(r"[A-Za-z]", sku) else "", sku if re.search(r"-\d{2}$", sku) else ""])
    _autosize(sheet, 3)
    return _save_book(book), len(rows)


def import_table(kind: str, data: bytes, filename: str) -> int:
    if kind not in TABLE_FILES:
        raise ValueError("Choose an active products or Dreamrobot table.")
    rows = read_rows(data, filename)
    if not rows or len(rows) < 2:
        raise ValueError("The table must contain a header and at least one data row.")
    headers = {str(value or "").strip().casefold() for value in rows[0]}
    if kind == "active" and not any("ndler_sku" in name for name in headers):
        raise ValueError("The active products table needs the Händler_SKU column.")
    if kind == "dreamrobot" and not {"a_seller_sku", "ek_preis"}.issubset(headers):
        raise ValueError("The Dreamrobot table needs a_seller_sku and ek_preis columns.")
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    for row in rows:
        sheet.append([value if isinstance(value, (str, int, float, bool, type(None))) else str(value) for value in row])
    target = DATA_DIR / TABLE_FILES[kind]
    temp = target.with_suffix(".tmp.xlsx")
    book.save(temp)
    temp.replace(target)
    return len(rows) - 1


def table_status() -> dict[str, dict[str, object]]:
    result = {}
    for kind, filename in TABLE_FILES.items():
        path = DATA_DIR / filename
        result[kind] = {"available": path.is_file(), "updated": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="minutes") if path.is_file() else None}
    return result


def _stored_rows(kind: str) -> list[list[object]]:
    path = DATA_DIR / TABLE_FILES[kind]
    if not path.is_file():
        raise ValueError(f"Import the {kind} table first.")
    return read_rows(path.read_bytes(), path.name)


def active_products_report() -> tuple[Path, int]:
    rows = _stored_rows("active")
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Active products")
    for row in rows:
        sheet.append(row)
    path = _save_book(book)
    return path, len(rows) - 1


def _column(headers: list[object], name: str) -> int:
    for index, value in enumerate(headers):
        if name in str(value or "").casefold():
            return index
    raise ValueError(f"Imported table is missing the {name} column.")


def database_prices_report() -> tuple[Path, int]:
    active = _stored_rows("active")
    dreamrobot = _stored_rows("dreamrobot")
    active_sku = _column(active[0], "ndler_sku")
    active_shipping = _column(active[0], "ndlerversandgruppe")
    active_spalte = _column(active[0], "spalte1")
    dream_sku = _column(dreamrobot[0], "a_seller_sku")
    dream_ek = _column(dreamrobot[0], "ek_preis")
    dream_vk = _column(dreamrobot[0], "vk_preis")
    products: dict[str, list[list[object]]] = {}
    for row in active[1:]:
        if len(row) > active_sku:
            products.setdefault(str(row[active_sku] or ""), []).append(row)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Database pricing")
    sheet.append(["Parent SKU", "Children SKU", "Label", "Wein/Feinkost", "Price Category", "Parent EK1er", "Children EKXer", "Faktor", "VK Children Final", "Transport Adjustment"])
    count = 0
    for row in dreamrobot[1:]:
        if len(row) <= max(dream_sku, dream_ek, dream_vk):
            continue
        sku = str(row[dream_sku] or "")
        if not sku.startswith("01-") or sku not in products:
            continue
        try:
            ek1er = (Decimal(str(row[dream_ek])) / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (ValueError, ArithmeticError):
            continue
        if ek1er == 0:
            continue
        for product in products[sku]:
            spalte = str(product[active_spalte] or "").strip().upper()
            shipping = str(product[active_shipping] or "").strip().lower()
            for label in BATCH_LABELS:
                result = calculate_pack(ek1er, label, "normal", spalte, database=True)
                raw = result["raw"]
                if raw < 50 and shipping in {"light", "standardvorlage amazon"}:
                    raw -= 3
                    note = "€3 DISCOUNT (VK < 50 & Light/Standard)"
                elif raw < 50:
                    note = "No Adjustment (VK < 50 & Prime/Versand)"
                else:
                    note = "No Adjustment (VK ≥ 50)"
                sheet.append([sku[3:], f"{result['prefix']}-{sku[3:]}", label, spalte, shipping.capitalize(),
                              f"{ek1er:.2f}", f"{result['ekxer']:.2f}", float(result["factor"]), f"{ending_99(raw):.2f}", note])
                count += 1
    return _save_book(book), count


PAGE = r'''
<h1>VS Database Project</h1>
<p class="hint">Calculate product prices and generate the same Excel reports as the VS Database project.</p>
<div class="db-actions">
  <button id="dbActive" class="db-green">PRODUKTET AKTIVE TE VS NE AMAZON</button>
  <button id="dbPrices" class="secondary">CALCULATE PRICES FROM DATABASE</button>
</div>
<div id="dbStatus" class="status" role="status"></div>
<div class="db-columns">
  <section class="panel"><h2>Calculate Pricing for One Product</h2>
    <div class="db-inputs"><div class="field"><label for="dbEk">Enter EK1er Price</label><input id="dbEk" type="number" min="0" step="0.01" value="6.3"></div>
    <fieldset><legend>Pricing Type</legend><label><input type="radio" name="dbType" value="normal" checked> Normal</label><label><input type="radio" name="dbType" value="feinkost"> Feinkost</label></fieldset></div>
    <div class="db-table"><table><thead><tr><th>Amazon</th><th>12er</th><th>6er</th><th>3er</th><th>4er</th><th>2er</th><th>1er</th></tr></thead><tbody id="dbManual"></tbody></table></div>
  </section>
  <section class="panel"><h2>Excel Format Example</h2><p class="hint">Columns: <strong>Produkt, EK1er</strong>. An optional third column, <strong>Spalte</strong> (W/F), controls Normal pricing. Other values use W.</p>
    <img class="db-example" src="/assets/cmimetVisando.png" alt="Excel example with Produkt and EK1er columns">
    <div class="field"><label for="dbExcel">Upload Excel File</label><input id="dbExcel" type="file" accept=".xlsx,.xlsm,.xls,.csv"></div>
    <div class="db-actions"><button id="dbNormal" class="secondary">Generate Normal Pricing</button><button id="dbFeinkost" class="db-cyan">Generate Feinkost Pricing</button></div>
  </section>
</div>
<section class="panel"><h2>Process Visando SKU File</h2><p class="hint">Upload CSV or Excel with <code>a_seller_sku</code> in the first column. The report marks SKUs without letters and SKUs ending in a hyphen and two digits.</p>
  <div class="field"><label for="dbSkuFile">Upload Visando SKU File</label><input id="dbSkuFile" type="file" accept=".csv,.txt,.xlsx,.xlsm,.xls"></div>
  <div class="actions"><button id="dbSku" class="db-green">Generate Processed SKU Report</button></div>
</section>
<section class="panel"><h2>Database source files</h2><p class="hint">Import exports of the original “PRODUKTET AKTIVE NE AMAZON” and “EXPORTI I PRODUKTEVE NGA DREAMROBOT” tables to generate the two database reports above. Files stay in the site's private data folder.</p>
  <div class="db-imports"><div class="field"><label for="dbActiveFile">Active products table</label><input id="dbActiveFile" type="file" accept=".xlsx,.xlsm,.xls,.csv"><button id="dbImportActive" class="mutedbtn">Import active products</button><small id="dbActiveState"></small></div>
  <div class="field"><label for="dbDreamFile">Dreamrobot products table</label><input id="dbDreamFile" type="file" accept=".xlsx,.xlsm,.xls,.csv"><button id="dbImportDream" class="mutedbtn">Import Dreamrobot products</button><small id="dbDreamState"></small></div></div>
</section>
<style>
.db-actions{display:flex;gap:12px;flex-wrap:wrap;margin:16px 0 24px}.db-green{background:#16bf87}.db-cyan{background:#16afbf}.db-columns{display:grid;grid-template-columns:1fr 1fr;gap:22px;margin-top:88px}.db-columns h2,.panel h2{font-size:21px;font-weight:500;color:#4b66db;margin:0 0 22px}.db-inputs{display:grid;grid-template-columns:1fr 1fr;gap:22px;margin-bottom:18px}.db-inputs fieldset{border:0;padding:0;margin:0}.db-inputs legend{font-size:13px;font-weight:750;margin-bottom:11px}.db-inputs fieldset label{font-weight:400;margin-right:12px}.db-inputs input[type=radio]{width:auto}.db-table{overflow:auto}.db-table th,.db-table td{text-align:center;min-width:75px}.db-table tbody tr:nth-child(odd){background:#f4f4f4}.db-table tbody th{white-space:nowrap}.db-example{max-width:100%;width:348px;height:auto;display:block;border:1px solid var(--line);border-radius:6px;margin:12px 0 24px}.db-imports{display:grid;grid-template-columns:1fr 1fr;gap:20px}.db-imports .field button{justify-self:start}.db-imports small{color:var(--muted)}@media(max-width:1100px){.db-columns{grid-template-columns:1fr;margin-top:32px}}@media(max-width:650px){.db-inputs,.db-imports{grid-template-columns:1fr}}
</style>
<script>
const dbPacks=[['12er',12,4.33,2.84],['6er',6,4.33,1.55],['3er',3,4.33,.85],['4er',4,4.33,1.55],['2er',2,3.60,.68],['1er',1,3.60,.42]];
function dbStatus(message,error=false){const e=document.getElementById('dbStatus');e.textContent=message;e.style.display='block';e.style.borderColor=error?'#b42318':'#0f766e'}
function dbRound99(raw){const whole=Math.floor(raw);return (raw-whole>=.5?whole+.99:whole-.01).toFixed(2)}
function dbUpdate(){const ek=Number(document.getElementById('dbEk').value);const type=document.querySelector('input[name="dbType"]:checked').value;const names=['EK1er','EKXer','Porto','Pick & Pack','Faktor','VK'];const lines=names.map((name,i)=>'<tr><th>'+name+'</th>'+dbPacks.map(([label,qty,porto,pack])=>{if(!(ek>0)&&![2,3].includes(i))return '<td>–</td>';const f=type==='feinkost'?(qty===12||qty===6||qty===3?1.53:1.48):(qty===12||qty===6?1.56:1.55);const values=[ek,ek*qty,porto,pack,f,dbRound99((ek*qty+porto+pack)*f)];return '<td>'+([4,5].includes(i)?(i===4?f.toFixed(2):values[i]):values[i].toFixed(2))+'</td>'}).join('')+'</tr>');document.getElementById('dbManual').innerHTML=lines.join('')}
async function dbDownload(url,file,extra={}){const form=new FormData();if(file)form.append('file',file);for(const [key,value] of Object.entries(extra))form.append(key,value);dbStatus('Generating report…');const response=await fetch(url,{method:file||Object.keys(extra).length?'POST':'GET',body:file||Object.keys(extra).length?form:undefined});if(!response.ok){let msg='Report failed';try{msg=(await response.json()).error||msg}catch{}throw Error(msg)}const blob=await response.blob();const disposition=response.headers.get('Content-Disposition')||'';const match=disposition.match(/filename="([^"]+)"/);const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download=match?match[1]:'VS_Database_Report.xlsx';a.click();setTimeout(()=>URL.revokeObjectURL(a.href),30000);dbStatus('Download started.')}
async function dbImport(kind,id){const file=document.getElementById(id).files[0];if(!file)throw Error('Choose a table export first.');const form=new FormData();form.append('file',file);form.append('kind',kind);dbStatus('Importing '+kind+' table…');const response=await fetch('/api/database/import',{method:'POST',body:form});const result=await response.json();if(!response.ok)throw Error(result.error||'Import failed');dbStatus('Imported '+result.rows+' rows.');dbLoadStatus()}
async function dbLoadStatus(){const r=await fetch('/api/database/status');if(!r.ok)return;const data=await r.json();for(const [kind,id] of [['active','dbActiveState'],['dreamrobot','dbDreamState']])document.getElementById(id).textContent=data[kind].available?'Imported '+data[kind].updated:'No table imported'}
function dbRun(job){Promise.resolve().then(job).catch(error=>dbStatus(error.message,true))}
document.getElementById('dbEk').addEventListener('input',dbUpdate);document.querySelectorAll('input[name="dbType"]').forEach(e=>e.addEventListener('change',dbUpdate));
document.getElementById('dbActive').addEventListener('click',()=>dbRun(()=>dbDownload('/api/database/active-products')));
document.getElementById('dbPrices').addEventListener('click',()=>dbRun(()=>dbDownload('/api/database/calculated-prices')));
document.getElementById('dbNormal').addEventListener('click',()=>dbRun(()=>dbDownload('/api/database/pricing',document.getElementById('dbExcel').files[0],{pricing_type:'normal'})));
document.getElementById('dbFeinkost').addEventListener('click',()=>dbRun(()=>dbDownload('/api/database/pricing',document.getElementById('dbExcel').files[0],{pricing_type:'feinkost'})));
document.getElementById('dbSku').addEventListener('click',()=>dbRun(()=>dbDownload('/api/database/sku',document.getElementById('dbSkuFile').files[0])));
document.getElementById('dbImportActive').addEventListener('click',()=>dbRun(()=>dbImport('active','dbActiveFile')));
document.getElementById('dbImportDream').addEventListener('click',()=>dbRun(()=>dbImport('dreamrobot','dbDreamFile')));
dbUpdate();dbLoadStatus().catch(()=>{});
</script>
'''
